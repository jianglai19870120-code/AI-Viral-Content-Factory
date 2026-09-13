"""标准化视频文案工作簿为逐视频、可追溯的 Markdown 原始资料。

原始工作簿永不改写。输出位于输入目录的 ``00_标准化源文档``，供后续
内容模块拆解使用；这里不做任何观点、痛点或误区判断。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable


ROOT = Path(os.environ.get("AI_TRAFFIC_FACTORY_ROOT") or Path(__file__).resolve().parents[3]).resolve()
DEFAULT_SOURCE_DIR = ROOT / "02_资产中心" / "01_输入库" / "04_视频文案-源文件（会员专享）"
OUTPUT_DIRNAME = "00_标准化源文档"
MANIFEST_JSONL = "manifest.jsonl"
MANIFEST_JSON = "manifest.json"

# 不依赖标题：同一账号下的原表序号才是 source_id 的稳定主键。
ACCOUNT_SLUGS = {
    "安先生": "anxiansheng",
    "白也": "baiye",
    "经纬": "jingwei-01",
    "经纬2号": "jingwei-02",
}

HEADER_ALIASES = {
    "序号": {"序号", "编号", "id"},
    "作品链接": {"作品链接", "视频链接", "链接", "url"},
    "作品标题": {"作品标题", "视频标题", "标题", "视频信息"},
    "点赞数": {"点赞数", "点赞", "赞"},
    "评论数": {"评论数", "评论"},
    "收藏数": {"收藏数", "收藏"},
    "转发数": {"转发数", "分享数", "转发"},
    "作品时长": {"作品时长", "视频时长", "时长"},
    "发布时间": {"发布时间", "发布时间", "发布日期", "发布时间"},
    "作品文案": {"作品文案", "视频文案", "文案", "全文", "原文"},
    "全文摘要": {"全文摘要", "摘要", "内容摘要"},
}


@dataclass(frozen=True)
class VideoRecord:
    source_id: str
    account: str
    account_slug: str
    source_file: str
    source_file_sha256: str
    source_sheet: str
    source_row: int
    source_serial: str
    title: str
    link: str
    source_fields: dict[str, Any]
    text_field: str
    full_text: str

    @property
    def full_text_sha256(self) -> str:
        return sha256_text(self.full_text)

    @property
    def relative_path(self) -> str:
        return f"{self.account_slug}/{self.source_id}.md"

    def manifest_entry(self) -> dict[str, Any]:
        return {
            "source_id": self.source_id,
            "source_type": "视频文案",
            "account": self.account,
            "source_file": self.source_file,
            "source_file_sha256": self.source_file_sha256,
            "source_sheet": self.source_sheet,
            "source_row": self.source_row,
            "source_serial": self.source_serial,
            "title": self.title,
            "link": self.link,
            "source_fields": self.source_fields,
            "text_field": self.text_field,
            "full_text_sha256": self.full_text_sha256,
            "relative_path": self.relative_path,
        }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="将视频文案工作簿标准化为逐视频 Markdown 源文档")
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR, help="包含原始工作簿的目录")
    parser.add_argument("--output-dir", type=Path, help="标准化输出目录，默认写入来源目录下的 00_标准化源文档")
    parser.add_argument("--apply", action="store_true", help="真正写入 Markdown 与来源清单；默认只预演")
    parser.add_argument("--check", action="store_true", help="校验已写入的来源清单、源文档和原始工作簿的一致性")
    return parser.parse_args()


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_text(value: str) -> str:
    return sha256_bytes(value.encode("utf-8"))


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return text.strip()


def normalize_header(value: Any) -> str:
    text = normalize_text(value).lower().replace("\u3000", "")
    return re.sub(r"[\s_\-:：()（）\[\]【】]", "", text)


def canonical_header(value: Any, index: int) -> str:
    normalized = normalize_header(value)
    for canonical, aliases in HEADER_ALIASES.items():
        if normalized in {normalize_header(alias) for alias in aliases}:
            return canonical
    raw = normalize_text(value)
    return raw if raw else f"字段{index + 1}"


def cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def is_legacy_ole(path: Path) -> bool:
    with path.open("rb") as handle:
        return handle.read(8) == bytes.fromhex("D0CF11E0A1B11AE1")


def read_workbook_rows(path: Path) -> list[tuple[str, int, tuple[Any, ...]]]:
    """读取 OOXML 或旧式 BIFF 工作簿；返回 (sheet, 1-based row, values)。"""
    if is_legacy_ole(path):
        try:
            import xlrd  # type: ignore
        except ImportError as exc:
            raise RuntimeError(
                f"{path.name} 是旧式 Excel 二进制文件；请安装 xlrd 后重试："
                "python -m pip install xlrd"
            ) from exc
        book = xlrd.open_workbook(path, on_demand=True)
        rows: list[tuple[str, int, tuple[Any, ...]]] = []
        for sheet in book.sheets():
            for row_index in range(sheet.nrows):
                values = tuple(cell_value(cell.value) for cell in sheet.row(row_index))
                rows.append((sheet.name, row_index + 1, values))
        book.release_resources()
        return rows

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - bundled runtime provides openpyxl
        raise RuntimeError("读取 .xlsx 需要 openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in workbook.worksheets:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            rows.append((sheet.title, row_index, tuple(cell_value(value) for value in row)))
    workbook.close()
    return rows


def header_score(row: Iterable[Any]) -> int:
    values = [normalize_header(value) for value in row]
    aliases = {normalize_header(alias) for group in HEADER_ALIASES.values() for alias in group}
    return sum(1 for value in values if value in aliases)


def find_headers(rows: list[tuple[str, int, tuple[Any, ...]]]) -> dict[str, tuple[int, tuple[Any, ...]]]:
    """为每个工作表找前 25 行内最像字段名的一行。"""
    by_sheet: dict[str, list[tuple[int, tuple[Any, ...]]]] = {}
    for sheet, row_number, values in rows:
        by_sheet.setdefault(sheet, []).append((row_number, values))
    headers: dict[str, tuple[int, tuple[Any, ...]]] = {}
    for sheet, candidates in by_sheet.items():
        candidates = candidates[:25]
        best = max(candidates, key=lambda item: header_score(item[1]), default=None)
        if best is None or header_score(best[1]) < 2:
            raise ValueError(f"{sheet} 未找到可识别的表头（前 25 行至少需要两个已知字段）")
        headers[sheet] = best
    return headers


def source_serial(value: Any, physical_row: int) -> str:
    text = normalize_text(value)
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        return f"{int(float(text)):06d}"
    if text:
        safe = re.sub(r"[^0-9A-Za-z_-]+", "-", text).strip("-")
        return safe[:40] or f"row-{physical_row:06d}"
    return f"row-{physical_row:06d}"


def account_metadata(source_path: Path) -> tuple[str, str]:
    account = source_path.stem.strip()
    slug = ACCOUNT_SLUGS.get(account)
    if not slug:
        slug = "account-" + sha256_text(account)[:12]
    return account, slug


def extract_records(source_dir: Path) -> list[VideoRecord]:
    files = sorted(path for path in source_dir.glob("*.xlsx") if path.is_file())
    if not files:
        raise FileNotFoundError(f"未在 {source_dir} 找到 .xlsx 工作簿")
    records: list[VideoRecord] = []
    used_ids: set[str] = set()
    for source_path in files:
        account, account_slug = account_metadata(source_path)
        source_hash = sha256_bytes(source_path.read_bytes())
        rows = read_workbook_rows(source_path)
        headers = find_headers(rows)
        for sheet, row_number, values in rows:
            header_row_number, header_values = headers[sheet]
            if row_number <= header_row_number:
                continue
            header_names = [canonical_header(value, index) for index, value in enumerate(header_values)]
            fields = {
                header_names[index]: cell_value(values[index]) if index < len(values) else ""
                for index in range(len(header_names))
            }
            full_text = normalize_text(fields.get("作品文案")) or normalize_text(fields.get("全文摘要"))
            if not full_text:
                continue
            serial = source_serial(fields.get("序号"), row_number)
            source_id = f"VS-{account_slug}-{serial}"
            if source_id in used_ids:
                raise ValueError(f"检测到重复 source_id：{source_id}；请先修复 {source_path.name} 的序号列")
            used_ids.add(source_id)
            title = normalize_text(fields.get("作品标题"))
            link = normalize_text(fields.get("作品链接"))
            source_fields = {key: value for key, value in fields.items() if key not in {"作品文案", "全文摘要"}}
            records.append(
                VideoRecord(
                    source_id=source_id,
                    account=account,
                    account_slug=account_slug,
                    source_file=source_path.name,
                    source_file_sha256=source_hash,
                    source_sheet=sheet,
                    source_row=row_number,
                    source_serial=serial,
                    title=title,
                    link=link,
                    source_fields=source_fields,
                    text_field="作品文案" if normalize_text(fields.get("作品文案")) else "全文摘要",
                    full_text=full_text,
                )
            )
    return records


def yaml_string(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def display_field_value(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return normalize_text(value).replace("\n", "<br>")


def render_record(record: VideoRecord) -> str:
    frontmatter = [
        "---",
        f"source_id: {yaml_string(record.source_id)}",
        'source_type: "视频文案"',
        f"account: {yaml_string(record.account)}",
        f"source_file: {yaml_string(record.source_file)}",
        f"source_file_sha256: {yaml_string(record.source_file_sha256)}",
        f"source_sheet: {yaml_string(record.source_sheet)}",
        f"source_row: {record.source_row}",
        f"source_serial: {yaml_string(record.source_serial)}",
        f"title: {yaml_string(record.title)}",
        f"link: {yaml_string(record.link)}",
        f"text_field: {yaml_string(record.text_field)}",
        f"full_text_sha256: {yaml_string(record.full_text_sha256)}",
        "---",
        "",
        f"# {record.source_id}",
        "",
        "## 来源字段",
        "",
        "| 字段 | 值 |",
        "|---|---|",
    ]
    frontmatter.extend(f"| {key} | {display_field_value(value)} |" for key, value in record.source_fields.items())
    frontmatter.extend(["", "## 全文", "", record.full_text, ""])
    return "\n".join(frontmatter)


def atomic_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(content, encoding="utf-8", newline="\n")
    temporary.replace(path)


def write_output(records: list[VideoRecord], output_dir: Path) -> None:
    for record in records:
        target = output_dir / record.relative_path
        content = render_record(record)
        if not target.exists() or target.read_text(encoding="utf-8") != content:
            atomic_write(target, content)

    entries = [record.manifest_entry() for record in records]
    jsonl = "".join(json.dumps(entry, ensure_ascii=False, sort_keys=True) + "\n" for entry in entries)
    atomic_write(output_dir / MANIFEST_JSONL, jsonl)
    atomic_write(output_dir / MANIFEST_JSON, json.dumps({"schema_version": "video-source-manifest-v1", "records": entries}, ensure_ascii=False, indent=2, sort_keys=True) + "\n")


def frontmatter_value(content: str, key: str) -> str:
    match = re.search(rf"(?m)^{re.escape(key)}:\s*(.+)$", content)
    if not match:
        return ""
    try:
        return str(json.loads(match.group(1)))
    except json.JSONDecodeError:
        return match.group(1).strip().strip('"')


def check_output(records: list[VideoRecord], output_dir: Path) -> list[str]:
    errors: list[str] = []
    manifest_path = output_dir / MANIFEST_JSONL
    if not manifest_path.exists():
        return [f"缺少来源清单：{manifest_path}"]
    manifest_records = [json.loads(line) for line in manifest_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    expected = {record.source_id: record for record in records}
    actual = {str(entry.get("source_id", "")): entry for entry in manifest_records}
    if len(actual) != len(manifest_records):
        errors.append("来源清单中存在重复 source_id")
    if set(actual) != set(expected):
        errors.append(f"来源清单 source_id 不一致：期望 {len(expected)} 条，实际 {len(actual)} 条")
    for source_id, record in expected.items():
        entry = actual.get(source_id)
        if entry is None:
            continue
        if entry.get("source_file_sha256") != record.source_file_sha256:
            errors.append(f"{source_id} 的原始文件 hash 不匹配")
        if entry.get("full_text_sha256") != record.full_text_sha256:
            errors.append(f"{source_id} 的全文 hash 不匹配")
        document = output_dir / record.relative_path
        if not document.exists():
            errors.append(f"缺少源文档：{document.relative_to(output_dir)}")
            continue
        content = document.read_text(encoding="utf-8")
        for key, expected_value in {
            "source_id": record.source_id,
            "source_file_sha256": record.source_file_sha256,
            "full_text_sha256": record.full_text_sha256,
        }.items():
            if frontmatter_value(content, key) != expected_value:
                errors.append(f"{source_id} 的 {key} 不匹配")
        marker = "\n## 全文\n\n"
        if marker not in content:
            errors.append(f"{source_id} 缺少全文正文段")
        elif sha256_text(content.split(marker, 1)[1].rstrip("\n")) != record.full_text_sha256:
            errors.append(f"{source_id} 的正文与工作簿全文不一致")
    return errors


def main() -> int:
    args = parse_args()
    source_dir = args.source_dir.resolve()
    output_dir = (args.output_dir or source_dir / OUTPUT_DIRNAME).resolve()
    records = extract_records(source_dir)
    by_account: dict[str, int] = {}
    for record in records:
        by_account[record.account] = by_account.get(record.account, 0) + 1
    print("读取工作簿：" + "、".join(sorted(path.name for path in source_dir.glob("*.xlsx"))))
    print("可读视频数：%d（%s）" % (len(records), "；".join(f"{name} {count}" for name, count in sorted(by_account.items()))))
    print(f"标准化输出目录：{output_dir}")
    if args.apply:
        write_output(records, output_dir)
        print("已写入逐视频源文档和来源清单。")
    elif not args.check:
        print("预演完成；加 --apply 才会写入。")
    if args.check:
        errors = check_output(records, output_dir)
        if errors:
            print("校验失败：")
            print("\n".join(f"- {error}" for error in errors))
            return 1
        print(f"校验通过：{len(records)} 条 source_id 均可回溯到原始工作簿、来源清单和全文。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
