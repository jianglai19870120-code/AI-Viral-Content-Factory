#!/usr/bin/env python
# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

try:
    from openpyxl import load_workbook
    from openpyxl import Workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("缺少依赖 openpyxl，无法读取标准 xlsx。") from exc

try:
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None

def bootstrap_brand_footer() -> None:
    candidate_roots: List[Path] = []

    if "--root" in sys.argv:
        try:
            root_arg = sys.argv[sys.argv.index("--root") + 1]
            candidate_roots.append(Path(root_arg).resolve())
        except (IndexError, OSError):
            pass

    cwd = Path.cwd().resolve()
    candidate_roots.extend([cwd, *cwd.parents])

    script_path = Path(__file__).resolve()
    candidate_roots.extend(script_path.parents)

    seen: set[str] = set()
    for root in candidate_roots:
        key = os.fspath(root)
        if key in seen:
            continue
        seen.add(key)
        # 全局唯一品牌尾注入口 = workflow/common.py::append_brand_footer
        wf = root / "workflow" / "common.py"
        if wf.exists():
            sys.path.insert(0, str(root))
            return


bootstrap_brand_footer()

try:
    from workflow.common import append_brand_footer
except ImportError:
    def append_brand_footer(text: str) -> str:
        return text


CATEGORIES = {
    "科学创业": "01_科学创业选题表.md",
    "能力成长": "02_能力成长选题表.md",
    "赚钱财富": "03_赚钱财富选题表.md",
    "个人IP": "04_个人IP选题表.md",
    "AI科技": "05_AI科技选题表.md",
    "其他类型": "99_其他类型选题表.md",
}

TABLE_HEADER = ["核心关键词", "选题", "原爆款元素", "博主名", "点赞数", "链接", "是否选中", "对标复刻拆解编号", "状态"]
LEGACY_V41_HEADER = TABLE_HEADER + ["文案结构"]
TOPIC_TABLE_GLOB = "*选题表.md"
HAND_INPUT_FILENAME = "00_手动输入选题表.xlsx"
XLS_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
XLSX_SIGNATURE = b"PK"
SOFT_TOPIC_LIMIT = 34
LONG_COMPLETE_LIMIT = 56
BAD_ENDING_RE = re.compile(r"(一|二|三|四|五|六|七|八|九|十|第|首先|其次|然后|因为|但是|而且|以及|如果|所以|例如|比如|其中|另外)$")
FRAGMENT_SUFFIX_RE = re.compile(
    r"(过后|之后|之间|的时候|阶段|故事|套路|系列|视角|原因|经验分享|想法|感觉|真相|秘诀|方向|自由|标准|阶段后)$"
)
OPENING_PUNCT = "([{【「『“'\""
CLOSING_PUNCT = ")]}】」』”'\""
EXPLANATION_STARTERS = [
    "AIDA代表",
    "很多人",
    "大多数人",
    "其实",
    "我觉得",
    "我前面",
    "我每次",
    "我经历的",
    "我做",
    "我用",
    "这里说的",
    "关键在于",
    "也就是",
    "区别只在于",
    "本质是",
    "记住",
    "相信我",
    "说实话",
    "如果你",
    "如果你的",
    "最好的",
    "真正的",
    "核心是",
    "如何发现",
    "把问题变成机会",
    "你的人生",
    "做知识博主能赚钱吗",
    "普通人如何赚钱",
    "用你的结果",
    "穷人用时间换钱",
    "你花在",
    "这条视频",
    "起点很低",
    "普通人",
    "短视频创作",
    "做短视频",
    "卖货的本质",
    "新人做",
    "真正重要的是",
    "第一步",
    "第二步",
    "第三步",
    "一、",
    "二、",
    "三、",
    "1、",
    "2、",
    "3、",
]
WEAK_SHORT_TOPICS = {
    "不露脸",
    "你的时间",
    "一边学习",
    "先走心",
    "换张地图",
}
GENERIC_LEAD_PHRASES = [
    "1条视频讲清楚",
    "一条视频讲清楚",
    "写给小白",
    "这条视频讲清楚",
]
MEANINGFUL_TAIL_PATTERNS = [
    r"如何变现$",
    r"怎么赚钱$",
    r"的区别$",
    r"有什么区别$",
    r"是什么$",
    r"怎么做$",
    r"怎么选$",
    r"怎么用$",
    r"个启发$",
    r"个思考$",
    r"种模式$",
    r"种方法$",
    r"条路线$",
    r"干抖音$",
]
EXPLANATORY_TAIL_MARKERS = [
    r"\s*提示词\s*\d*\s*[：:]",
    r"\s*\|\s*",
    r"\s*我是一个",
    r"\s*请帮我",
    r"\s*以下是",
    r"\s*首先",
    r"\s*其次",
    r"\s*然后",
    r"\s*第一步",
    r"\s*第二步",
    r"\s*第三步",
    r"\s*[①②③④⑤]",
]
TAIL_STARTERS = (
    "普通人", "很多", "大多数", "我是", "我想", "我用", "我做", "你", "如果", "当",
    "因为", "这是", "这个", "这种", "首先", "其次", "然后", "第一步", "第二步", "第三步",
    "进入", "不要", "不是", "每天", "分享", "选择", "想", "看完", "投资", "现在", "做",
    "从", "真正", "了解", "短视频", "学习", "小龙", "道",
)

WEALTH_STRONG_KEYWORDS = [
    "赚钱", "挣钱", "搞钱", "收入", "副业", "单干", "生意", "第一桶金",
    "年入", "月入", "变富", "财务自由", "被动收入", "资产", "现金流",
    "变现", "值钱", "赚到钱", "养活自己", "赚点小钱", "赚到", "财富",
    "第一笔", "100w", "100万", "生产资料", "致富", "富翁",
]
PERSONAL_IP_STRONG_KEYWORDS = [
    "自媒体", "账号", "涨粉", "起号", "卖课", "课程", "知识博主", "女粉",
    "短视频", "小红书", "抖音", "私域", "内容产品", "个人品牌", "内容获客",
]
CAREER_GROWTH_KEYWORDS = [
    "找工作", "面试", "空窗期", "工作能力", "转行", "职业能力", "社交指南",
]


@dataclass
class TopicRow:
    keyword: str
    original_topic: str
    original_elements: str
    primary_bucket: str
    blogger: str
    likes: str
    link: str
    optimized_topic: str
    is_selected: str = ""
    benchmark_case_id: str = ""
    status: str = ""

    def key(self) -> Tuple[str, str]:
        if self.link:
            return ("link", self.link)
        return ("topic", f"{self.blogger}::{self.original_topic}")

    def to_cells(self) -> List[str]:
        return [
            self.keyword,
            self.original_topic,
            self.original_elements,
            self.blogger,
            self.likes,
            self.link,
            self.is_selected,
            self.benchmark_case_id,
            self.status,
        ]


def normalize_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def markdown_escape(value: str) -> str:
    value = normalize_cell(value)
    value = value.replace("\\", "\\\\").replace("|", "\\|")
    value = value.replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")
    return value


def split_markdown_row(line: str) -> List[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    cells: List[str] = []
    current: List[str] = []
    escaped = False
    for char in text:
        if char == "\\" and not escaped:
            escaped = True
            current.append(char)
            continue
        if char == "|" and not escaped:
            cells.append("".join(current).strip().replace("\\|", "|"))
            current = []
        else:
            current.append(char)
        escaped = False
    cells.append("".join(current).strip().replace("\\|", "|"))
    return cells


def is_separator_row(cells: List[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells if cell.strip())


def compact_text(raw: str) -> str:
    text = normalize_cell(raw)
    text = text.split("#", 1)[0]
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"(展开|收起|复制链接|DOU\+小助手)$", "", text, flags=re.I).strip()
    text = re.sub(r"([\s|｜/\-—]+抖音)$", "", text, flags=re.I).strip()
    return text


def split_sentences(text: str) -> List[str]:
    parts = re.split(r"(?<=[。！？?!])\s*|…{2,}|\.\.\.+", text)
    return [part.strip() for part in parts if part.strip()]


def normalize_topic_candidate(text: str) -> str:
    candidate = re.sub(r"\s+", " ", text).strip(" ，,；;：:。.!?！？、-…")
    return candidate.strip()


def has_explanatory_tail(text: str) -> bool:
    return bool(any(re.search(pattern, text) for pattern in EXPLANATORY_TAIL_MARKERS))


def cut_explanatory_tail(text: str) -> Tuple[str, bool, str]:
    candidate = normalize_topic_candidate(text)
    if not candidate:
        return "", False, ""
    if candidate.startswith("问道 "):
        return "问道", True, "两字题眼加解释串，只保留题眼"
    first_punct = re.search(r"[。！？?!]", candidate)
    if first_punct and 4 <= first_punct.start() <= 44 and len(candidate) - first_punct.end() >= 18:
        head = normalize_topic_candidate(candidate[: first_punct.start()])
        if head and not ends_like_fragment(head):
            nested_head, nested_cut, _ = cut_explanatory_tail(head)
            if nested_cut and nested_head:
                return nested_head, True, "首句加解释串，只保留首句题眼"
            return head, True, "首句加解释串，只保留首句题眼"
    for pattern in EXPLANATORY_TAIL_MARKERS:
        match = re.search(pattern, candidate)
        if match and 4 <= match.start() <= 80:
            head = normalize_topic_candidate(candidate[: match.start()])
            if head and not ends_like_fragment(head):
                nested_colon = re.match(r"^(?P<head>[^ ]{6,36}[：:][^ ]{4,28})\s+(?P<tail>.+)$", head)
                if nested_colon:
                    nested_colon_head = normalize_topic_candidate(nested_colon.group("head"))
                    nested_colon_tail = normalize_topic_candidate(nested_colon.group("tail"))
                    if nested_colon_tail.startswith(TAIL_STARTERS) or len(nested_colon_tail) >= 18:
                        return nested_colon_head, True, "切除提示词、竖线或口播解释尾巴后，再保留冒号标题"
                nested = re.match(r"^(?P<head>[^。！？?!：:，,；;]{4,44}?)\s+(?P<tail>.+)$", head)
                if nested:
                    nested_head = normalize_topic_candidate(nested.group("head"))
                    nested_tail = normalize_topic_candidate(nested.group("tail"))
                    if nested_tail.startswith(TAIL_STARTERS) or len(nested_tail) >= 18:
                        return nested_head, True, "切除提示词、竖线或口播解释尾巴后，再保留题眼"
                return head, True, "切除提示词、竖线或口播解释尾巴"
    comma_question = re.match(r"^[^，,]{2,12}[，,](?P<head>(如何|怎么|为什么|能不能|有没有)[^，,。！？?!\s]{4,30})(?:\s|$)", candidate)
    if comma_question:
        head = normalize_topic_candidate(comma_question.group("head"))
        if head and not ends_like_fragment(head):
            return head, True, "去掉前置标签，保留真正选题"
    colon_title = re.match(r"^(?P<head>[^ ]{6,36}[：:][^ ]{4,28})\s+(?P<tail>.+)$", candidate)
    if colon_title:
        head = normalize_topic_candidate(colon_title.group("head"))
        tail = normalize_topic_candidate(colon_title.group("tail"))
        if tail.startswith(TAIL_STARTERS) or len(tail) >= 18:
            return head, True, "冒号标题加案例解释，只保留完整标题"
    loose_eye = re.match(r"^(?P<head>[^ ]{4,28})\s+(?P<tail>.+)$", candidate)
    if loose_eye:
        head = normalize_topic_candidate(loose_eye.group("head"))
        tail = normalize_topic_candidate(loose_eye.group("tail"))
        if head and tail and not ends_like_fragment(head) and (tail.startswith(TAIL_STARTERS) or len(tail) >= 30):
            return head, True, "题眼加解释串，只保留题眼"
    topic_eye = re.match(r"^(?P<head>[^。！？?!：:，,；;]{4,44}?)\s+(?P<tail>.+)$", candidate)
    if topic_eye:
        head = normalize_topic_candidate(topic_eye.group("head"))
        tail = normalize_topic_candidate(topic_eye.group("tail"))
        if (
            head
            and tail
            and not ends_like_fragment(head)
            and (
                tail.startswith(TAIL_STARTERS)
                or len(tail) >= 26
                or re.search(r"(的本质|的方法|的核心|是什么|为什么|怎么|如何)", tail)
            )
            and not any(re.search(pattern, tail, flags=re.I) for pattern in MEANINGFUL_TAIL_PATTERNS)
        ):
            return head, True, "题眼加解释串，只保留题眼"
    return candidate, False, ""


def dedupe_repeated_topic(text: str) -> str:
    normalized = normalize_topic_candidate(text)
    if not normalized:
        return normalized
    parts = [part.strip() for part in re.split(r"[。！？?!]", normalized) if part.strip()]
    if len(parts) >= 2 and parts[0] == parts[1]:
        return parts[0]
    if len(normalized) % 2 == 0:
        half = len(normalized) // 2
        left = normalized[:half].strip()
        right = normalized[half:].strip()
        if left and left == right:
            return left
    return normalized


def has_unmatched_opening(text: str) -> bool:
    stack: List[str] = []
    pairs = dict(zip(CLOSING_PUNCT, OPENING_PUNCT))
    symmetric_pairs = {'"', "'"}
    for char in text:
        if char in symmetric_pairs:
            if stack and stack[-1] == char:
                stack.pop()
            else:
                stack.append(char)
            continue
        if char in OPENING_PUNCT:
            stack.append(char)
        elif char in CLOSING_PUNCT:
            if stack and stack[-1] == pairs[char]:
                stack.pop()
    return bool(stack)


def ends_like_fragment(text: str) -> bool:
    if not text:
        return True
    if BAD_ENDING_RE.search(text):
        return True
    if FRAGMENT_SUFFIX_RE.search(text):
        return True
    if has_unmatched_opening(text):
        return True
    return bool(re.search(r"[：:（(「『“\"]$", text))


def looks_like_weak_short_topic(text: str) -> bool:
    if text in WEAK_SHORT_TOPICS:
        return True
    if text.endswith("？") or text.endswith("?"):
        return False
    if re.search(r"(做事|成事|资产|投资|成长|方法|路径|系统|模型|策略|财富|创业|赚钱|变现|销售|工作流|能力|认知|习惯)$", text):
        return False
    if len(text) <= 6 and not re.search(r"(如何|怎么|为什么|是否|会不会|能不能|是不是)", text):
        return True
    return False


def is_complete_topic(text: str) -> bool:
    candidate = normalize_topic_candidate(text)
    if not candidate:
        return False
    if ends_like_fragment(candidate):
        return False
    if looks_like_weak_short_topic(candidate):
        return False
    return True


def split_clauses(text: str) -> List[str]:
    return [part.strip() for part in re.split(r"(…{2,}|\.\.\.+|[，；;。！？?!])", text) if part and part.strip()]


def split_space_phrases(text: str) -> List[str]:
    return [part.strip() for part in text.split(" ") if part.strip()]


def looks_like_hook_phrase(text: str) -> bool:
    if not text:
        return False
    candidate = normalize_topic_candidate(text)
    if len(candidate) < 5 or len(candidate) > 28:
        return False
    if ends_like_fragment(candidate):
        return False
    return bool(re.search(r"(怎么|如何|为什么|能不能|有没有|什么时候|多少|区别|秘诀|方法|收入来源|赚钱|买房|起号|翻盘|变现|自由|启发|指南|攻略|真相|认知|策略)", candidate))


def trim_label_tail(text: str) -> str:
    raw_candidate = re.sub(r"\s+", " ", str(text)).strip()
    candidate = normalize_topic_candidate(raw_candidate)
    if not candidate:
        return ""
    explicit_followup = re.match(
        r"^(?P<head>.+?)\s+(?P<followup>(做|靠|用|学|买|选).{0,14}(赚钱吗|怎么赚钱|能赚钱吗|值不值|靠不靠谱|有没有必要|好不好|行不行))$",
        candidate,
    )
    if explicit_followup:
        head = normalize_topic_candidate(explicit_followup.group("head"))
        if len(head) >= 6 and is_complete_topic(head):
            return head
    raw_phrases = [part.strip() for part in raw_candidate.split(" ") if part.strip()]
    phrases = [normalize_topic_candidate(part) for part in raw_phrases]
    if len(phrases) >= 2:
        first = normalize_topic_candidate(phrases[0])
        raw_remainder = " ".join(raw_phrases[1:])
        remainder = normalize_topic_candidate(raw_remainder)
        if first and remainder:
            if ("？" in raw_remainder or "?" in raw_remainder) and len(first) >= 6:
                return first
            if looks_like_hook_phrase(remainder) and len(first) >= 6:
                return first
            if remainder.startswith(("你", "我", "如果", "当", "因为", "这是", "也就是", "普通人")) and len(first) >= 6:
                return first
    return candidate


def extract_structured_head(text: str) -> str:
    candidate = normalize_topic_candidate(text)
    if not candidate:
        return ""

    generic_lead = re.match(r"^(?P<lead>[^：:]{2,16})[：:\s]+(?P<body>.+)$", candidate)
    if generic_lead:
        lead = normalize_topic_candidate(generic_lead.group("lead"))
        body = normalize_topic_candidate(generic_lead.group("body"))
        if lead in GENERIC_LEAD_PHRASES and body and is_complete_topic(body):
            return body

    numbered_with_followup = re.match(
        r"^(?P<prefix>.+?)\s+[一二三四五六七八九十]+、(?P<section>[^。！？?!]{4,40}?)\s+(?P<followup>[^。！？?!]{4,24})$",
        candidate,
    )
    if numbered_with_followup:
        prefix = normalize_topic_candidate(numbered_with_followup.group("prefix"))
        section = normalize_topic_candidate(numbered_with_followup.group("section"))
        followup = normalize_topic_candidate(numbered_with_followup.group("followup"))
        if prefix and section and looks_like_hook_phrase(followup):
            combined = normalize_topic_candidate(f"{prefix}：{section}")
            if is_complete_topic(combined):
                return combined

    numbered = re.match(r"^(?P<prefix>.+?)\s+[一二三四五六七八九十]+、(?P<section>[^。！？?!]{4,30})", candidate)
    if numbered:
        prefix = normalize_topic_candidate(numbered.group("prefix"))
        section = trim_label_tail(numbered.group("section"))
        if prefix and section:
            combined = normalize_topic_candidate(f"{prefix}：{section}")
            if is_complete_topic(combined):
                return combined

    decimal_numbered = re.match(r"^(?P<prefix>.+?)\s+0\.(?P<section>[^。！？?!]{4,30})", candidate)
    if decimal_numbered:
        prefix = normalize_topic_candidate(decimal_numbered.group("prefix"))
        section = trim_label_tail(decimal_numbered.group("section"))
        if prefix and section:
            combined = normalize_topic_candidate(f"{prefix}：{section}")
            if is_complete_topic(combined):
                return combined

    lead_with_list = re.match(r"^(?P<intro>.+?)\s+(?P<body>[^。！？?!]{6,32}[：:])\s*[0-9一二三四五六七八九十]+[、.]", candidate)
    if lead_with_list:
        body = normalize_topic_candidate(lead_with_list.group("body").rstrip("：:"))
        if is_complete_topic(body):
            return body

    prefix_with_numbered_list = re.match(
        r"^(?P<head>[^。！？?!]{6,34}?)\s+[0-9一二三四五六七八九十]+[、:：]",
        candidate,
    )
    if prefix_with_numbered_list:
        head = normalize_topic_candidate(prefix_with_numbered_list.group("head"))
        if is_complete_topic(head):
            return head

    phrases = split_space_phrases(candidate)
    if len(phrases) >= 2:
        first = normalize_topic_candidate(phrases[0])
        second = normalize_topic_candidate(" ".join(phrases[1:]))
        if first in GENERIC_LEAD_PHRASES and second and is_complete_topic(second):
            return second
        if looks_like_hook_phrase(first) and second.startswith(("如果", "你", "普通人", "一个人", "我", "这", "做", "想", "当")):
            return first
        if first and second and re.match(r"^[0-9一二三四五六七八九十]+[、.]", second):
            return first

    return candidate


def first_sentence(text: str) -> str:
    sentences = split_sentences(text)
    return sentences[0] if sentences else text.strip()


def truncate_to_boundary(text: str, limit: int) -> str:
    if len(text) <= limit:
        return normalize_topic_candidate(text)
    boundary = max(
        text.rfind("，", 0, limit + 1),
        text.rfind("；", 0, limit + 1),
        text.rfind("。", 0, limit + 1),
        text.rfind(" ", 0, limit + 1),
        text.rfind("：", 0, limit + 1),
        text.rfind("、", 0, limit + 1),
    )
    if boundary >= 8:
        return normalize_topic_candidate(text[:boundary])
    return normalize_topic_candidate(text[:limit])


def trim_explanatory_tail(text: str) -> str:
    candidate = extract_structured_head(dedupe_repeated_topic(text))
    if not candidate:
        return ""
    direct_head, did_cut, _ = cut_explanatory_tail(candidate)
    if did_cut and direct_head:
        return direct_head
    space_parts = candidate.split(" ", 1)
    if len(space_parts) == 2:
        first = normalize_topic_candidate(space_parts[0])
        second = normalize_topic_candidate(space_parts[1])
        if first.endswith("就是") and second:
            merged = normalize_topic_candidate(f"{first} {second}")
            if len(merged) <= LONG_COMPLETE_LIMIT and is_complete_topic(merged):
                return merged
        if (
            4 <= len(first) <= 30
            and is_complete_topic(first)
            and second.startswith(("当", "如果", "人会", "我一直", "你周围", "投资", "现在", "\"", "“", "‘"))
        ):
            return first
    sentence_parts = re.split(r"[。！？?!]\s+", candidate, maxsplit=1)
    if len(sentence_parts) == 2:
        head = normalize_topic_candidate(sentence_parts[0])
        if 4 <= len(head) <= 30 and is_complete_topic(head):
            return head
    sentence_head = re.match(r"^(?P<head>[^。！？?!]{4,30}[。！？?!])\s+(?P<tail>.+)$", candidate)
    if sentence_head:
        head = normalize_topic_candidate(sentence_head.group("head"))
        if is_complete_topic(head):
            return head
    starter_head = re.match(
        r"^(?P<head>[^。！？?!]{4,30})\s+(?P<tail>(当|如果|人会|我一直|你周围|投资|现在|\"|“|‘).+)$",
        candidate,
    )
    if starter_head:
        head = normalize_topic_candidate(starter_head.group("head"))
        if is_complete_topic(head):
            return head
    judgement_pair = re.match(r"^(?P<head>[^。！？?!]{4,24}就是)\s+(?P<tail>[^。！？?!]{2,24})$", candidate)
    if judgement_pair:
        merged = normalize_topic_candidate(f"{judgement_pair.group('head')} {judgement_pair.group('tail')}")
        if is_complete_topic(merged):
            return merged
    if any(re.search(pattern, candidate, flags=re.I) for pattern in MEANINGFUL_TAIL_PATTERNS):
        return candidate
    if len(candidate) <= LONG_COMPLETE_LIMIT and is_complete_topic(candidate):
        return candidate
    lead_phrases = split_space_phrases(candidate)
    if len(lead_phrases) >= 2:
        first = normalize_topic_candidate(lead_phrases[0])
        second = normalize_topic_candidate(lead_phrases[1])
        if first.endswith("就是"):
            return candidate
        if (
            4 <= len(first) <= 28
            and is_complete_topic(first)
            and not any(re.search(pattern, second, flags=re.I) for pattern in MEANINGFUL_TAIL_PATTERNS)
            and second.startswith(("很多", "最好的", "普通人", "借", "学会", "把", "如何", "核心", "与", "就是", "真正", "关键", "第一步", "第二步", "第三步", "人会", "你", "我", "我们", "现在", "如果", "当", "不是", "想", "\"", "“", "‘"))
        ):
            return first
    starter_head = re.match(
        r"^(?P<head>.+?)\s+(?P<tail>(最好的|真正的|核心是|如何发现|把问题变成机会|你的人生|做知识博主能赚钱吗|普通人如何赚钱|用你的结果|穷人用时间换钱).+)$",
        candidate,
    )
    if starter_head:
        head = normalize_topic_candidate(starter_head.group("head"))
        if len(head) >= 6 and is_complete_topic(head):
            return head
    numbered_followup_head = re.match(
        r"^(?P<prefix>.+?)\s+[一二三四五六七八九十]+、(?P<section>[^。！？?!]{4,40}?)\s+(?P<followup>(做|靠|用|学|买|选).{0,14}(赚钱吗|怎么赚钱|能赚钱吗|值不值|靠不靠谱|有没有必要|好不好|行不行))(?:[。！？?!]|$)",
        candidate,
    )
    if numbered_followup_head:
        prefix = normalize_topic_candidate(numbered_followup_head.group("prefix"))
        section = normalize_topic_candidate(numbered_followup_head.group("section"))
        if prefix and section:
            combined = normalize_topic_candidate(f"{prefix}：{section}")
            if is_complete_topic(combined):
                return combined
    question_cut = max(candidate.find("？"), candidate.find("?"))
    if question_cut >= 8 and question_cut < len(candidate) - 1:
        head = normalize_topic_candidate(candidate[: question_cut + 1])
        if is_complete_topic(head):
            return head
    for starter in EXPLANATION_STARTERS:
        for token in (f" {starter}", f"，{starter}", f"：{starter}", f"。{starter}"):
            position = candidate.find(token)
            if position >= 8:
                head = normalize_topic_candidate(candidate[:position])
                if is_complete_topic(head):
                    return head
    return candidate


def compress_if_overlong(text: str) -> Tuple[str, bool, str]:
    normalized = dedupe_repeated_topic(text)
    if not normalized:
        return "", False, ""
    direct_head, did_cut, cut_reason = cut_explanatory_tail(normalized)
    if did_cut and direct_head:
        return direct_head, True, cut_reason
    if len(normalized) <= SOFT_TOPIC_LIMIT and is_complete_topic(normalized):
        return normalized, False, ""
    if len(normalized) <= LONG_COMPLETE_LIMIT and is_complete_topic(normalized) and not has_explanatory_tail(normalized):
        return normalized, False, ""

    sentences = split_sentences(normalized)
    sentence = trim_explanatory_tail(first_sentence(normalized))
    if sentence != normalized and is_complete_topic(sentence) and len(sentence) <= LONG_COMPLETE_LIMIT:
        return sentence, True, "多句内容，只保留首个完整主张句"

    clauses = split_clauses(sentence)
    candidate = ""
    for piece in clauses:
        merged = normalize_topic_candidate(candidate + piece)
        if not merged:
            continue
        candidate = merged
        if len(candidate) <= SOFT_TOPIC_LIMIT and is_complete_topic(candidate):
            return candidate, True, "压缩口播解释串，保留最小完整句"

    phrases = split_space_phrases(sentence)
    if len(phrases) > 1:
        candidate = ""
        for part in phrases:
            merged = normalize_topic_candidate((candidate + " " + part).strip())
            if merged:
                candidate = merged
            if len(candidate) <= LONG_COMPLETE_LIMIT and is_complete_topic(candidate):
                trimmed = trim_explanatory_tail(candidate)
                return trimmed, True, "压缩空格串联解释，保留首个完整主张句"

    if is_complete_topic(sentence) and len(sentence) <= LONG_COMPLETE_LIMIT:
        return trim_explanatory_tail(sentence), True, "去掉后续展开，只保留首个完整主张句"

    running = ""
    best = sentence if sentence else normalized
    best_reason = "标题过长，保留首个完整分句"
    for piece in clauses:
        merged = normalize_topic_candidate(running + piece)
        if not merged:
            continue
        running = merged
        if len(running) <= LONG_COMPLETE_LIMIT and is_complete_topic(running):
            best = running
            best_reason = "标题过长，压缩为完整分句"

    if is_complete_topic(best):
        return trim_explanatory_tail(best), True, best_reason

    fallback = normalize_topic_candidate(sentence or normalized)
    bounded = truncate_to_boundary(fallback, LONG_COMPLETE_LIMIT)
    if bounded and is_complete_topic(bounded):
        return trim_explanatory_tail(bounded), True, "标题过长，按边界截到可用主张句"
    if fallback and not ends_like_fragment(fallback):
        return fallback, True, "标题过长，回退到最短可用主张句"

    return normalized, False, ""


def extract_primary_statement(text: str) -> Tuple[str, bool, str]:
    normalized = dedupe_repeated_topic(text)
    if not normalized:
        return "", False, ""
    direct_head, did_cut, cut_reason = cut_explanatory_tail(normalized)
    if did_cut and direct_head:
        return direct_head, True, cut_reason
    trimmed = trim_explanatory_tail(normalized)
    if trimmed != normalized and is_complete_topic(trimmed) and len(trimmed) <= LONG_COMPLETE_LIMIT:
        return trimmed, True, "压缩解释串，保留最小完整主张句"
    if is_complete_topic(normalized) and len(normalized) <= LONG_COMPLETE_LIMIT:
        return normalized, False, ""
    return compress_if_overlong(normalized)


def clean_topic(raw: str) -> Tuple[str, bool, str, str]:
    text = compact_text(raw)
    if not text:
        return "", False, "", ""
    topic, shortened, reason = extract_primary_statement(text)
    return topic, shortened, text, reason


def score_keywords(text: str, keywords: Iterable[str]) -> int:
    lower_text = text.lower()
    score = 0
    for kw in keywords:
        normalized_kw = kw.lower()
        if re.search(r"[a-z0-9]", normalized_kw):
            pattern = re.compile(rf"(?<![a-z0-9]){re.escape(normalized_kw)}(?![a-z0-9])", re.I)
            if pattern.search(lower_text):
                score += 1
        elif normalized_kw in lower_text:
            score += 1
    return score


def contains_any(text: str, keywords: Iterable[str]) -> bool:
    lower_text = text.lower()
    return any(keyword.lower() in lower_text for keyword in keywords)


def has_personal_ip_strong_signal(text: str) -> bool:
    return contains_any(text, PERSONAL_IP_STRONG_KEYWORDS)


def has_wealth_strong_signal(text: str) -> bool:
    return contains_any(text, WEALTH_STRONG_KEYWORDS)


def has_pure_career_growth_signal(text: str) -> bool:
    return contains_any(text, CAREER_GROWTH_KEYWORDS)


def infer_suspected_category(topic: str) -> Tuple[str, str]:
    lower_text = topic.lower()

    ai_keywords = [
        "ai", "agent", "gpt", "claude", "karpathy", "deepseek", "lex fridman", "prompt",
        "自动化", "工作流", "模型", "网站", "数据中心", "软件", "gui", "cli", "编码", "编程",
    ]
    ip_keywords = [
        "自媒体", "账号", "涨粉", "起号", "卖课", "课程", "知识博主", "赛道", "大主播",
        "内容", "女粉", "流量", "私域", "个人ip", "个人品牌", "选题", "文案",
    ]
    growth_keywords = [
        "成事", "做事", "做成", "执行力", "专注", "时间管理", "上班", "打工", "不上班",
        "空窗期", "面试", "方向", "喜欢", "感兴趣", "翻盘", "找工作", "社交", "人生",
        "个人提升", "成长", "习惯", "解决问题", "信息源", "生命周期",
    ]
    wealth_keywords = [
        "赚钱", "收入", "副业", "搞钱", "财富", "投资", "理财", "保险", "买房", "租房",
        "房租", "月供", "中介", "谈价", "房子", "房价", "金价", "关税", "贸易战",
        "经济", "医保", "翻倍", "变现",
    ]
    startup_keywords = [
        "创业", "商业", "生意", "利润", "客户", "产品", "增长", "成交", "销售", "过滤客户",
        "高端客户", "经营", "团队", "公司",
    ]

    if has_personal_ip_strong_signal(topic):
        return "个人IP", "含 自媒体 / 账号 / 涨粉 / 卖课 / 内容获客 等强个人IP语义"
    if has_wealth_strong_signal(topic):
        return "赚钱财富", "含 赚钱 / 收入 / 副业 / 生意 / 资产 / 财富结果 等强财富语义"
    if any(keyword in lower_text for keyword in ai_keywords):
        return "AI科技", "含 AI / 模型 / 自动化 / 科技产品 等语义"
    if any(keyword in lower_text for keyword in ip_keywords):
        return "个人IP", "含 自媒体 / 账号 / 涨粉 / 卖课 / 赛道 等语义"
    if any(keyword in lower_text for keyword in wealth_keywords):
        return "赚钱财富", "含 收入 / 财富 / 房产 / 宏观经济 / 变现 等语义"
    if any(keyword in lower_text for keyword in growth_keywords):
        return "能力成长", "含 成事 / 做事 / 职场处境 / 个人提升 等语义"
    if any(keyword in lower_text for keyword in startup_keywords):
        return "科学创业", "含 创业 / 生意 / 利润 / 客户 / 产品经营 等语义"
    return "待细化主题", "当前仍未命中稳定规则，需要继续补充通用分类口径"


def infer_other_display_category(topic: str, blogger: str) -> Tuple[str, str]:
    lower_text = f"{blogger} {topic}".lower()

    if any(keyword in lower_text for keyword in ["room tour", "四房两卫", "串串房", "居住展示", "入住", "一镜到底"]):
        return "居住展示记录", "主题聚焦居住空间、入住展示或看房后体验"

    if any(keyword in lower_text for keyword in ["城市", "base", "成都", "北京", "换城市", "定居", "生活方式"]):
        return "城市迁移与居住选择", "主题聚焦城市迁移、定居选择或生活方式切换"

    if any(keyword in lower_text for keyword in ["狗狗", "宠物", "大理", "玉龙雪山", "带狗旅行"]):
        return "宠物旅行", "主题聚焦宠物陪伴下的旅行或出行经验"

    if any(keyword in lower_text for keyword in ["采访", "对话", "播客截选", "访谈视频", "创始人对话", "截选了一小段"]):
        return "人物访谈片段", "主题聚焦人物访谈、播客截选或对话内容"

    if any(keyword in lower_text for keyword in ["年会", "线下", "厦门年会", "长沙线下", "活动现场"]):
        return "线下活动记录", "主题聚焦年会、线下见面或活动现场记录"

    if any(keyword in lower_text for keyword in ["我上央视了", "近况通报", "上央视"]):
        return "个人事件播报", "主题聚焦个人事件、近况或阶段性播报"

    if any(keyword in lower_text for keyword in ["信息密度", "刷视频状态", "内容观察", "编导在线答疑", "转发", "小作文"]):
        return "内容方法观察", "主题聚焦内容表达、平台观看状态或内容方法观察"

    if any(keyword in lower_text for keyword in ["科技现象", "前沿", "agi", "华为", "小米", "美国制造业", "豆包"]):
        return "科技现象观察", "主题聚焦技术前沿、科技公司或科技现象观察"

    return "抽象认知观点", "主题是泛认知、抽象判断或高概念观点，暂不稳定落入正式五类"


def classify_topic(topic: str) -> Tuple[str, str]:
    text = topic
    topic_text = topic

    ai_keywords = [
        "ai", "人工智能", "大模型", "模型", "agent", "智能体", "自动化", "编程", "代码",
        "claude", "gemini", "gpt", "openai", "cursor", "vibe coding", "芯片", "机器人",
        "工作流", "提示词", "token", "lex fridman", "karpathy", "deepseek",
    ]
    ip_keywords = [
        "自媒体", "短视频", "内容", "账号", "涨粉", "流量", "私域", "个人ip", "ip",
        "定位", "粉丝", "直播", "小红书", "抖音", "视频号", "公众号", "成交", "获客",
        "爆款", "文案", "选题", "剪辑", "起号", "卖课", "课程", "知识博主", "女粉",
        "内容产品", "个人品牌",
    ]
    money_keywords = [
        "赚钱", "收入", "副业", "财富", "现金流", "变现", "赚到", "赚", "钱",
        "商业机会", "普通人", "搞钱", "财富自由", "咨询", "接单", "投资", "理财",
        "生意", "创业赚钱",
    ]
    growth_keywords = [
        "学习", "认知", "效率", "行动", "行动力", "决策", "表达", "习惯", "复盘",
        "自律", "成长", "能力", "职业", "思考", "心态", "选择", "拖延", "焦虑",
        "知识", "做事", "现卖现学", "闭环", "执行力", "专注", "专注力", "时间管理",
        "上班", "打工", "空窗期", "面试", "人生方向", "自由职业", "成事", "小事",
        "做成", "超级个体", "个人提升",
    ]
    startup_keywords = [
        "创业", "商业模式", "公司", "企业", "产品", "组织", "增长", "战略", "融资",
        "管理", "团队", "企业服务", "创新", "商业化", "老板", "ceo", "客户",
    ]

    topic_scores = {
        "AI科技": score_keywords(topic_text, ai_keywords),
        "个人IP": score_keywords(topic_text, ip_keywords),
        "赚钱财富": score_keywords(topic_text, money_keywords),
        "能力成长": score_keywords(topic_text, growth_keywords),
        "科学创业": score_keywords(topic_text, startup_keywords),
    }
    scores = {
        name: topic_scores[name]
        for name in topic_scores
    }

    lower_text = text.lower()

    # 主承诺边界：个人IP强方法优先；明确财富结果优先于成长；AI 只有在工具/科技是主承诺时优先。
    if scores["个人IP"] >= 1 and has_personal_ip_strong_signal(text):
        return "个人IP", "主承诺围绕内容账号、流量或个人品牌"

    if has_wealth_strong_signal(text) and not has_personal_ip_strong_signal(text):
        return "赚钱财富", "主承诺围绕赚钱、收入、生意、资产或财富结果"

    if scores["AI科技"] >= 1 and any(k in lower_text for k in ["ai", "人工智能", "大模型", "模型", "agent", "智能体", "自动化", "claude", "gpt", "gemini", "cursor", "vibe coding", "工作流", "token", "提示词"]):
        if scores["个人IP"] == 0 and not has_wealth_strong_signal(text):
            return "AI科技", "主承诺围绕 AI、工具、模型或自动化应用"

    if any(k in text for k in ["买房", "租房", "房租", "月供", "房价", "房子", "中介", "谈价", "医保", "保险", "金价", "关税", "贸易战", "通货收缩"]):
        return "赚钱财富", "主承诺围绕财富决策、房产或宏观经济"

    if any(k in text for k in ["赛道", "大主播", "知识博主", "卖课", "起号", "涨粉", "女粉", "自媒体怎么", "做内容", "内容团队"]):
        return "个人IP", "主承诺围绕做账号、做内容或个人IP经营"

    if any(k in text for k in ["成事", "做事", "做成", "上班", "打工", "不上班", "空窗期", "面试", "找工作", "社交指南", "人生", "方向", "执行力", "专注力", "时间管理", "经验产品化", "信息源", "解决问题", "感兴趣", "喜欢的事", "翻盘"]):
        if has_wealth_strong_signal(text) and not has_pure_career_growth_signal(text):
            return "赚钱财富", "同时涉及成长/职业与财富，但主承诺是赚钱、收入或生意结果"
        return "能力成长", "主承诺围绕个人成长、做事能力或职场处境"

    if any(k in lower_text for k in ["karpathy", "deepseek", "lex fridman", "gavin baker", "网站", "数据中心", "软件", "gui", "cli"]) and scores["赚钱财富"] == 0:
        return "AI科技", "主承诺围绕 AI 观察、科技产品或技术范式"

    if any(k in text for k in ["利润", "高端客户", "过滤客户", "低端客户", "生意越小", "规模做小", "做利润"]):
        return "科学创业", "主承诺围绕利润、客户筛选或经营策略"

    if scores["能力成长"] >= 1 and any(k in text for k in ["执行力", "专注", "专注力", "时间管理", "上班", "打工", "空窗期", "面试", "人生方向", "自律", "拖延", "心态", "成事", "做事", "做成", "小事", "超级个体", "个人提升"]):
        if has_wealth_strong_signal(text) and not has_pure_career_growth_signal(text):
            return "赚钱财富", "成长词与财富词冲突时，以收入、生意或资产承诺为主"
        if scores["个人IP"] == 0:
            return "能力成长", "主承诺围绕个人成长、职业处境或能力提升"

    if scores["赚钱财富"] >= 2 and any(k in text for k in ["赚钱", "副业", "收入", "赚到", "变现", "接单", "财富", "生意"]):
        return "赚钱财富", "主承诺围绕赚钱、收入或变现"

    if topic_scores["AI科技"] > 0 and scores["AI科技"] >= 2 and scores["个人IP"] == 0 and scores["能力成长"] == 0:
        return "AI科技", "主承诺围绕 AI、工具、模型或科技产品"

    if topic_scores["能力成长"] > 0 and scores["个人IP"] == 0:
        return "能力成长", "主承诺围绕能力、认知或自我成长"

    if scores["科学创业"] >= 2 and scores["赚钱财富"] == 0 and scores["个人IP"] == 0:
        return "科学创业", "主承诺围绕创业、公司经营或商业化"

    if scores["能力成长"] >= 2 and scores["赚钱财富"] == 0:
        return "能力成长", "主承诺围绕能力、认知或自我成长"

    if scores["个人IP"] >= 1 and scores["赚钱财富"] >= 1:
        if any(k in lower_text for k in ["自媒体", "账号", "涨粉", "起号", "卖课", "课程", "知识博主", "女粉"]):
            return "个人IP", "同时涉及变现与账号，但主承诺更偏个人IP"

    if scores["能力成长"] >= 1 and scores["赚钱财富"] >= 1:
        if has_wealth_strong_signal(text) and not has_pure_career_growth_signal(text):
            return "赚钱财富", "同时涉及赚钱与成长，但主承诺更偏收入、生意或财富结果"
        if any(k in text for k in ["空窗期", "面试", "专注力", "时间管理", "执行力", "成事", "做事", "做成", "小事", "超级个体"]):
            return "能力成长", "同时涉及赚钱与成长，但主承诺更偏个人成长"

    best_category, best_score = max(scores.items(), key=lambda item: item[1])
    if best_score <= 0:
        return "其他类型", "信息不足或不属于当前五类"

    tied = [name for name, score in scores.items() if score == best_score]
    if len(tied) > 1:
        if "AI科技" in tied and topic_scores["AI科技"] > 0 and scores["赚钱财富"] == 0 and scores["个人IP"] == 0:
            return "AI科技", ""
        return "其他类型", "多类关键词冲突，需人工复核：" + "、".join(tied)

    return best_category, ""


CHATGPT_TOPIC_OVERRIDES: Dict[str, Tuple[str, str]] = {
    "人一旦掌握了赚钱思维，钱就会源源不断地流向你": ("赚钱思维", "普通人一旦掌握这3种赚钱思维，赚钱会越来越容易"),
    "25年下半年怎样赚点小钱？": ("副业赚钱", "普通人2025年下半年，怎么低成本赚到第一笔副业收入？"),
    "如何在一个加速分流的时代往上走，而不是往下滑？": ("阶层跃迁", "普通人如何在加速分流的时代，避免一路向下滑？"),
    "房租2800，月供3000，买房还是租房？": ("房产选择", "房租2800、月供3000，普通人到底该买房还是租房？"),
    "中美二次贸易战，战的到底是什么？": ("贸易博弈", "中美二次贸易战，真正争的根本不是关税"),
    "房租2800，月供3000，怎么选？": ("房产选择", "房租2800、月供3000，哪一种选择的长期成本更低？"),
    "为什么很多人1万都掏不出来，却觉得100万很少？": ("财富认知", "为什么越拿不出1万的人，越觉得100万不算多？"),
    "买房和租房哪个更划算？": ("房产选择", "普通人买房还是租房，哪一种长期成本更低？"),
    "AI的正确提效方式": ("AI提效", "多数人用AI提效都用反了：真正有效的3种方式"),
    "逾期别恐慌，积极面对是良方": ("债务处理", "花呗、借呗逾期的人，最不该做的3件事"),
    "经济发展背后的秘密，未来3-5年明确的方向": ("经济趋势", "未来3-5年，普通人最应该关注的3个经济方向"),
    "1个人，1个月，30万粉丝，自媒体如何作弊？": ("涨粉方法", "1个人、1个月涨粉30万：我如何用“作弊思维”做自媒体？"),
    "一条视频讲透怎么做自媒体": ("自媒体", "一条视频讲透：普通人怎么从0开始做自媒体"),
    "不上班，一个人在家怎样赚点小钱？": ("居家副业", "普通人不上班、没团队，怎么在家赚到第一笔钱？"),
    "普通人的金融世界观：容易被忽略的三点": ("金融认知", "普通人最容易忽略的3条金融规则"),
    "26年做自媒体确定性的3个趋势": ("自媒趋势", "2026年做自媒体，普通人能抓住的3个确定性趋势"),
    "人生开挂，10倍加速的成长系统，让我在字节年薪翻10倍的方法": ("职场成长", "我在字节把年薪翻10倍，只靠这套成长系统"),
    "人生第一笔20万，应该先买车还是先买房？": ("资产配置", "人生第一笔20万，普通人该先买车还是先买房？"),
    "如何把握你的生命周期": ("人生周期", "普通人最容易错过的3个生命周期红利"),
    "中美关税博弈，谁赢了？": ("关税博弈", "中美关税博弈，真正的赢家可能不是中美"),
    "普通人搞钱第一课：原始积累": ("原始积累", "普通人搞钱第一课：先完成原始积累"),
    "26年入局自媒体，你必须要知道的真相": ("自媒避坑", "2026年入局自媒体，普通人最容易踩的3个坑"),
    "人生开挂：我在字节年薪翻10倍的方法": ("职场成长", "我在字节如何把年薪翻10倍：普通人也能复制的3步"),
    "成年人社交指南：一条视频教你礼尚往来": ("人情往来", "成年人最容易吃亏的3种礼尚往来，一条视频讲透"),
    "中美对抗为什么外国都很安静？": ("国际博弈", "中美对抗越激烈，为什么多数国家反而更安静？"),
    "如果通货收缩来了，谁会最吃亏？": ("通货收缩", "通货收缩真来了，普通人里谁会最先吃亏？"),
    "人生靠什么翻盘：伴侣、贵人和成就自己": ("人生翻盘", "普通人想翻盘，伴侣、贵人和自己，哪个最重要？"),
    "找工作的核心思路，匹配风口与锚定趋势": ("求职选择", "普通人找工作，最重要的不是能力，而是站对趋势"),
    "不上班，一个人在家，怎么赚点小钱？": ("居家副业", "没资源的普通人，不上班怎么在家赚到第一笔钱？"),
    "油车和电车，2025年该怎么选？": ("汽车选择", "2025年普通人买车，油车和电车谁的长期成本更低？"),
    "如何赚到人生的第一个100w？": ("百万路径", "普通人如何从0赚到人生第一个100万？"),
    "如何赚到人生的第一个100万？": ("百万路径", "普通人如何从0赚到人生第一个100万？"),
    "自媒体没有风口，做IP没有捷径 大道至简，复杂的都是错的": ("IP方法", "自媒体没有风口，做IP没有捷径：越复杂的方法越可能是错的"),
    "如何让你的时间变得值钱？": ("时间价值", "普通人如何用3步，让自己的时间越来越值钱？"),
    "金价为什么会涨，为什么会跌？": ("黄金投资", "普通人买黄金前，必须看懂金价涨跌的3个原因"),
    "什么人一打眼，就知道工作能力很强": ("工作能力", "真正工作能力强的人，一眼就能看出的5个特征"),
    "2026记住存钱的三个核心关键词": ("存钱理财", "2026年普通人存钱，只记住这3个核心关键词"),
    "富豪们为什么要隐瞒长寿真相？": ("长寿认知", "富豪为什么不愿公开真正的长寿方法？"),
    "成年人社交指南，搞定各种社交关系": ("社交关系", "成年人最容易搞砸的3种社交关系"),
    "国家欠钱还不上了，会怎样？": ("国家债务", "国家也会欠钱不还？结局可能比个人逾期更严重"),
    "为什么美国制裁华为，但没有制裁小米": ("科技博弈", "美国为什么制裁华为，却放过了小米？"),
    "为什么我建议大家一定要多刷小红书": ("平台认知", "多数人劝你少刷小红书，我却建议你一定要多刷"),
    "如何杀出底层？逆风局的6条翻盘法则": ("底层翻盘", "普通人如何杀出底层：逆风局的6条翻盘法则"),
    "赚钱根本不靠能力，全靠心力，3招立刻拥有强大心力": ("心力成长", "赚钱不靠能力，靠心力：普通人立刻变强的3招"),
    "烂工作和失业哪个更痛苦？": ("工作选择", "烂工作和失业，普通人到底该选哪个？"),
    "年入百万是什么感觉？": ("百万收入", "年入百万的人，真实生活到底是什么样？"),
    "男性消费最真实的样子：不声张，但认准了，就会一直买下去": ("男性消费", "男性消费最反常的3个真相：不声张，但认准就一直买"),
    "普通家庭如何逆袭，鸡娃不如鸡自己": ("家庭逆袭", "普通家庭想逆袭，最该鸡的不是孩子，而是自己"),
    "一个人做自媒体没团队没资源，有没有什么靠谱的变现方式？": ("自媒变现", "没团队、没资源的普通人，做自媒体最靠谱的3种变现方式"),
    "普通人的第一桶金为什么会越来越难？": ("第一桶金", "普通人的第一桶金越来越难，不是因为机会变少了"),
    "AI的正确打开方式": ("AI应用", "90%的人都用错了AI：真正的打开方式只有3步"),
    "为什么DeepSeek爆火后我们想到的是裁员而不是上三休四？": ("AI就业", "DeepSeek越强，为什么我们先想到裁员，而不是上三休四？"),
    "不付费，AI就会变傻吗？": ("AI付费", "不付费的AI真的会变傻吗？我测试了3个版本"),
    "越平易近人越吃亏，反着做事更容易成": ("处世原则", "越平易近人的人越容易吃亏：反着做的3条规则"),
    "1个人，1年，100万粉丝，自媒体如何作弊？": ("涨粉方法", "1个人、1年涨粉100万：自媒体到底怎么“作弊”？"),
    "26年做自媒体的3个确定性趋势": ("自媒趋势", "2026年做自媒体，普通人能抓住的3个确定性趋势"),
    "中美关税博弈，谁能赢？": ("关税博弈", "中美关税博弈，最后真正买单的是谁？"),
    "如果工厂里全是机器人，还存在剩余价值吗？": ("剩余价值", "如果工厂全是机器人，老板还能赚走剩余价值吗？"),
    "真正的自律，是做好这两件事情": ("自律成长", "真正的自律，不是逼自己坚持，而是做好这两件事"),
    "做了四个月自媒体以后，我看见了3个红利": ("自媒红利", "做了4个月自媒体，我看见普通人还能抓住的3个红利"),
    "未来5-10年，每个人都有一次财富跃迁的机会": ("财富跃迁", "未来5-10年，普通人最可能实现财富跃迁的3次机会"),
    "第一性原理：先回到本质，再做出选择": ("第一原理", "第一性原理：多数人做选择，第一步就做反了"),
    "1周涨粉10万纯干货分享": ("涨粉方法", "1周涨粉10万：没团队的普通人做对了哪3件事？"),
    "有医保就看病不愁？医保到底应该怎么算": ("医保认知", "有医保就看病不愁？普通人最容易算错的3笔账"),
    "哪些房子在回暖，哪些还在下行？": ("房产趋势", "房价回暖不是普涨：哪3类房子在涨，哪3类还在跌？"),
    "做出属于你自己的AI内容团队": ("AI团队", "普通人如何用3步，做出自己的AI内容团队？"),
    "美国制造业能否回流成功？": ("制造回流", "美国制造业回流，最难解决的其实不是工厂"),
    "普通人搞错变富的顺序，真的会越忙越穷": ("变富顺序", "普通人变富最容易搞错的3步：越忙反而越穷"),
    "豆包的正确打开方式：离前沿更近一步": ("豆包应用", "90%的人都低估了豆包：真正的打开方式是这3步"),
    "居家自由职业1年，涨粉百万的时间管理心得": ("时间管理", "居家自由职业1年涨粉百万：我最有效的3条时间管理原则"),
    "AI时代的核心差距从来都不是会不会用，而是你怎么用": ("AI认知", "AI时代真正拉开差距的，不是会不会用，而是怎么用"),
}


KEYWORD_TO_BUCKET: Dict[str, str] = {
    "创业起步": "科学创业",
    "客户成交": "科学创业",
    "产品定价": "科学创业",
    "产品设计": "科学创业",
    "商业模式": "科学创业",
    "利润模型": "科学创业",
    "增长策略": "科学创业",
    "一人公司": "科学创业",
    "市场验证": "科学创业",
    "副业赚钱": "赚钱财富",
    "第一桶金": "赚钱财富",
    "财富认知": "赚钱财富",
    "资产配置": "赚钱财富",
    "房产选择": "赚钱财富",
    "房产趋势": "赚钱财富",
    "金融认知": "赚钱财富",
    "原始积累": "赚钱财富",
    "赚钱思维": "赚钱财富",
    "自媒体起号": "个人IP",
    "涨粉方法": "个人IP",
    "内容定位": "个人IP",
    "IP变现": "个人IP",
    "账号避坑": "个人IP",
    "平台认知": "个人IP",
    "自媒趋势": "个人IP",
    "内容拍摄": "个人IP",
    "职场成长": "能力成长",
    "工作选择": "能力成长",
    "求职选择": "能力成长",
    "执行力": "能力成长",
    "认知成长": "能力成长",
    "时间管理": "能力成长",
    "人生翻盘": "能力成长",
    "社交关系": "能力成长",
    "人情往来": "能力成长",
    "读书方法": "能力成长",
    "自我修炼": "能力成长",
    "学习方法": "能力成长",
    "目标管理": "能力成长",
    "AI应用": "AI科技",
    "AI提效": "AI科技",
    "AI团队": "AI科技",
    "AI就业": "AI科技",
    "模型工具": "AI科技",
    "自动化工作流": "AI科技",
    "AI认知": "AI科技",
    "AI付费": "AI科技",
    "豆包应用": "AI科技",
    "国际博弈": "其他类型",
    "经济趋势": "其他类型",
    "贸易博弈": "其他类型",
    "关税博弈": "其他类型",
    "科技博弈": "其他类型",
    "制造回流": "其他类型",
    "通货收缩": "赚钱财富",
    "国家债务": "赚钱财富",
    "黄金投资": "赚钱财富",
    "医保认知": "赚钱财富",
    "男性消费": "其他类型",
    "家庭逆袭": "能力成长",
    "人生周期": "能力成长",
    "心力成长": "能力成长",
    "自律成长": "能力成长",
    "城市生活": "其他类型",
    "宠物旅行": "其他类型",
    "居住体验": "其他类型",
    "读书推荐": "其他类型",
    "线下活动": "其他类型",
    "人物访谈": "其他类型",
    "生活方式": "其他类型",
    "生活观察": "其他类型",
    "个人表达": "能力成长",
    "机会识别": "能力成长",
    "规则思维": "能力成长",
    "工作复利": "能力成长",
    "内容传播": "个人IP",
    "带货视频": "个人IP",
    "演讲表达": "能力成长",
    "阶层跃迁": "能力成长",
    "能源选择": "赚钱财富",
    "长寿认知": "赚钱财富",
    "AI趋势": "AI科技",
    "科技产品": "AI科技",
    "社群避坑": "能力成长",
    "资源交换": "能力成长",
    "综合观察": "其他类型",
}

KEYWORD_RULES: List[Tuple[str, List[str]]] = [
    ("市场验证", ["持续存在的需求", "市场验证", "判断一个生意", "能不能做", "需求", "试错", "实际问题"]),
    ("客户成交", ["客户买单", "客户付钱", "高端客户", "低端客户", "成交", "销售", "购买", "付钱", "收费越高", "尊重你", "框架越清晰", "主张越锋利", "不卖而卖", "卖货的本质"]),
    ("创业起步", ["创业起步", "起步阶段", "创业第一步", "从 0 到 1", "从0到1", "0成本创业", "10分钟创业法", "低成本创业", "创业方法", "不要轻易创业"]),
    ("一人公司", ["一人公司", "你自己就是一家公司", "你自己就可以是一家公司", "一个人公司"]),
    ("产品定价", ["定价", "价格", "收费", "卖多少钱"]),
    ("产品设计", ["产品设计", "产品太多", "产品太弱", "卖点", "场景+理念", "买产品", "普通商品"]),
    ("利润模型", ["利润", "规模做小", "利润做大", "别做规模", "回本", "现金流"]),
    ("增长策略", ["增长", "10倍增长", "2倍", "用户增长"]),
    ("商业模式", ["商业模式", "开源模式", "商业闭环", "经营模型", "运营团队", "Github开源", "Skill产品"]),
    ("AI团队", ["AI内容团队", "AI团队"]),
    ("AI提效", ["提效", "效率"]),
    ("AI就业", ["裁员", "上三休四", "AI就业"]),
    ("AI付费", ["付费", "不付费"]),
    ("豆包应用", ["豆包"]),
    ("模型工具", ["DeepSeek", "Claude", "GPT", "模型", "大模型", "Trae", "Windsurf", "Token"]),
    ("自动化工作流", ["工作流", "自动化", "Agent", "Skill", "Codex"]),
    ("AI应用", ["AI", "人工智能"]),
    ("涨粉方法", ["涨粉", "粉丝", "大主播", "播放量", "万粉", "百万粉丝"]),
    ("自媒体起号", ["起号", "冷启动", "新手怎么入门", "普通人能不能做自媒体", "怎么做自媒体"]),
    ("内容定位", ["内容定位", "赛道", "选题", "推流", "内容表达", "做内容", "文案"]),
    ("IP变现", ["内容变现", "IP变现", "挣钱方式", "赚钱方式", "卖课", "课程", "私域", "个人品牌"]),
    ("账号避坑", ["自媒体谣言", "常犯", "自嗨", "割韭菜", "靠谱的老师", "高粉低变现"]),
    ("平台认知", ["小红书", "抖音", "平台", "推流机制"]),
    ("自媒趋势", ["自媒体", "趋势", "红利", "风口"]),
    ("内容拍摄", ["拍摄", "光线", "画面", "编导", "镜头", "视频演示", "王者拍摄法"]),
    ("内容传播", ["转发", "小作文", "刷视频", "观众", "传播", "分享有", "视频", "博主", "上央视"]),
    ("带货视频", ["带货视频", "跨境带货"]),
    ("房产选择", ["买房", "租房", "房租", "月供", "买车还是买房"]),
    ("房产趋势", ["房价", "房子", "回暖", "下行", "二手房"]),
    ("资产配置", ["20万", "第一笔", "资产配置", "资产"]),
    ("副业赚钱", ["副业", "赚点小钱", "不上班", "在家", "自由职业", "单干"]),
    ("第一桶金", ["第一桶金", "第一笔钱"]),
    ("原始积累", ["原始积累", "生产资料"]),
    ("金融认知", ["金融", "医保", "黄金", "债务", "逾期", "存钱", "理财", "金价", "保险"]),
    ("财富认知", ["财富", "变富", "财务自由", "100万", "百万", "钱", "越忙越穷", "财富跃迁"]),
    ("赚钱思维", ["赚钱思维", "赚钱真相", "赚钱", "搞钱", "挣钱", "收入"]),
    ("职场成长", ["字节", "年薪", "成长系统", "工作能力", "职场", "上班很强"]),
    ("求职选择", ["找工作", "面试", "转行", "空窗期"]),
    ("工作选择", ["烂工作", "失业", "上班", "打工", "不上班"]),
    ("执行力", ["执行力", "行动", "做成", "做事", "成事", "专注"]),
    ("认知成长", ["认知", "第一性原理", "开窍", "思维", "世界观"]),
    ("时间管理", ["时间管理", "时间", "自律"]),
    ("读书方法", ["读完", "读书", "临摹", "自学", "看书", "字帖"]),
    ("学习方法", ["学什么", "学习", "技能", "终极技能", "两项技能", "进入一个领域", "快速进入"]),
    ("社群避坑", ["私董会"]),
    ("资源交换", ["拿你有的换你想要的", "大佬主动找你", "无许可门徒"]),
    ("目标管理", ["目标", "计划", "长期主义", "完成一件", "10倍经验"]),
    ("自我修炼", ["内耗", "看见自己", "障碍", "我不行", "变得强大", "脑袋里的", "先做完", "从简单处开始"]),
    ("个人表达", ["内向", "慢热", "上场", "演讲", "灵感", "写点什么", "枯燥点的视频", "不是英语老师", "钥匙"]),
    ("机会识别", ["机会", "规则之外", "金矿", "洞见", "前辈", "听得进", "帮助你成功"]),
    ("规则思维", ["作弊", "智径", "r/k策略", "策略", "方法有三个核心优势", "三种价值", "4个价值"]),
    ("工作复利", ["有积累的工作", "没有积累的工作", "最好的工作", "印钞机", "工作其实", "一次构建终身受益"]),
    ("演讲表达", ["就这么演讲", "大神上身"]),
    ("社交关系", ["社交", "关系", "伴侣", "贵人"]),
    ("人情往来", ["礼尚往来", "人情"]),
    ("人生翻盘", ["翻盘", "逆袭", "底层", "杀出底层"]),
    ("家庭逆袭", ["家庭", "鸡娃"]),
    ("人生周期", ["生命周期"]),
    ("心力成长", ["心力"]),
    ("自律成长", ["自律"]),
    ("经济趋势", ["经济", "通货", "制造业", "贸易", "关税", "中美"]),
    ("国际博弈", ["中美对抗", "外国都很安静"]),
    ("科技博弈", ["华为", "小米", "制裁"]),
    ("AI趋势", ["AGI", "硅谷巨头", "巨头慌了"]),
    ("科技产品", ["钢铁侠盔甲", "咖啡助理", "小玩具"]),
    ("能源选择", ["油车", "电车"]),
    ("长寿认知", ["长寿"]),
    ("阶层跃迁", ["阶层跃迁", "加速分流", "往上走", "往下滑"]),
    ("男性消费", ["男性消费"]),
    ("城市生活", ["换城市", "成都", "城市", "Base", "大理", "厦门"]),
    ("宠物旅行", ["狗狗", "小狗", "宠物"]),
    ("居住体验", ["住宿", "Room tour", "串串房", "四房两卫"]),
    ("读书推荐", ["推荐几本书", "采访视频", "价值心法"]),
    ("线下活动", ["年会", "线下"]),
    ("人物访谈", ["对话", "创始人", "大哥"]),
    ("生活方式", ["理想的生活", "爱好变成了作品", "活法"]),
]


def normalize_override_key(text: str) -> str:
    return normalize_topic_candidate(text).replace("—", "-").replace("，", ",")


def detect_original_elements(topic: str) -> str:
    elements: List[str] = []
    if re.search(r"(\d+|[一二三四五六七八九十]+)\s*(个|步|招|条|种|分钟|小时|天|周|月|年|万|元|块|倍|%|％)", topic) or any(k in topic for k in ["零基础", "低成本", "一条视频", "捷径", "不辞职", "不用", "没团队", "没资源", "低门槛"]):
        elements.append("成本")
    if any(k in topic for k in ["普通人", "新人", "宝妈", "职场", "成年人", "男性", "女生", "家庭", "中年", "30到50岁", "没资源", "没团队", "小白", "月薪"]):
        elements.append("人群")
    if any(k in topic for k in ["奇葩", "离谱", "全是机器人", "作弊", "从来没人"]):
        elements.append("奇葩")
    if any(k in topic for k in ["最差", "难用", "难吃", "烂工作", "割韭菜", "低端客户", "逾期", "裁员", "失业", "踩坑", "坑", "谣言"]):
        elements.append("难吃难用差")
    if re.search(r"(不是.+而是|不是.+是|却|反着|反而|越.+越|不靠|不如|放弃|不存在|别|为什么.+但|不是.*靠)", topic):
        elements.append("反向操作")
    if any(k in topic for k in ["怀旧", "童年", "小时候", "80后", "90后", "过去"]):
        elements.append("怀旧")
    if any(k in topic for k in ["吸引女粉", "女粉", "男性消费", "伴侣", "恋爱", "吸引力", "长寿"]):
        elements.append("荷尔蒙")
    if any(k in topic for k in ["最大", "最全", "最强", "最贵", "第一", "头部", "大主播", "富豪", "字节", "华为", "小米", "DeepSeek", "全网", "百万", "10倍"]):
        elements.append("头牌")
    return "＋".join(dict.fromkeys(elements)) if elements else "无明显元素"


def infer_core_keyword(topic: str, primary_bucket: str = "") -> str:
    override = CHATGPT_TOPIC_OVERRIDES.get(normalize_override_key(topic))
    if override:
        return override[0]
    for keyword, needles in KEYWORD_RULES:
        if any(needle in topic for needle in needles):
            return keyword
    suspected_category, _ = infer_suspected_category(topic)
    if suspected_category == "科学创业":
        return "商业模式"
    if suspected_category == "赚钱财富":
        return "赚钱思维"
    if suspected_category == "个人IP":
        return "内容定位"
    if suspected_category == "AI科技":
        return "AI应用"
    if suspected_category == "能力成长":
        return "认知成长"
    if primary_bucket == "科学创业":
        return "商业模式"
    if primary_bucket == "能力成长":
        return "认知成长"
    if primary_bucket == "赚钱财富":
        return "赚钱思维"
    if primary_bucket == "个人IP":
        return "内容定位"
    if primary_bucket == "AI科技":
        return "AI应用"
    return "综合观察"


def keyword_to_primary_bucket(keyword: str, fallback: str) -> str:
    if keyword in KEYWORD_TO_BUCKET:
        return KEYWORD_TO_BUCKET[keyword]
    if fallback in CATEGORIES:
        return fallback
    return "其他类型"


def has_problematic_topic_residue(topic: str) -> bool:
    if len(topic) > LONG_COMPLETE_LIMIT:
        return True
    return bool(
        "|" in topic
        or "提示词" in topic
        or re.search(r"\s(我是一个|请帮我|以下是|首先|其次|然后|第一步|第二步|第三步|①|②|③)", topic)
    )


def has_strong_topic_hook(topic: str) -> bool:
    return bool(re.search(
        r"(为什么|怎么|如何|能不能|有没有|到底|不是|而是|反而|却|越.+越|最|第一|"
        r"\d+\s*(分钟|小时|天|周|月|年|万|元|块|倍|步|个|种|条|%|％))",
        topic,
    ))


def ensure_enhanced_when_no_element(base: str, optimized: str, keyword: str, elements: str) -> str:
    optimized = re.sub(r"\s+", " ", optimized).strip(" ，,；;：:、-…")
    if elements != "无明显元素":
        return optimized
    if optimized and optimized != base:
        return optimized
    if keyword in {"综合观察", "生活观察", "生活方式"}:
        return f"{base}，背后真正值得警惕的是什么？"
    return f"{base}，最容易忽略的3个关键是什么？"


def optimize_topic(topic: str, keyword: str, elements: str) -> str:
    override = CHATGPT_TOPIC_OVERRIDES.get(normalize_override_key(topic))
    if override:
        return override[1]
    base = normalize_topic_candidate(topic)
    exact_rewrites = {
        "10分钟创业法": "10分钟能不能跑通一门小生意？",
        "高端客户付钱，是因为不付的代价更贵": "高端客户为什么愿意付钱？因为不付的代价更贵",
        "一人公司起步阶段，90%的人注定失败": "一人公司起步阶段，为什么90%的人注定失败？",
        "创业的方法 创业方法论 找到持续存在的需求": "创业方法论：怎么找到持续存在的真实需求？",
        "如何快速进入一个领域": "如何快速进入一个领域：先抓住20%核心知识",
    }
    if base in exact_rewrites:
        return ensure_enhanced_when_no_element(base, exact_rewrites[base], keyword, elements)
    if elements != "无明显元素" and has_strong_topic_hook(base):
        return base
    if keyword == "客户成交":
        return f"{base}，客户真正买单的原因是什么？"
    if keyword == "创业起步":
        return f"{base}，第一步到底该怎么做？"
    if keyword == "一人公司":
        return f"{base}，一个人怎么跑通？"
    if keyword == "市场验证":
        return f"{base}，怎么用低成本验证？"
    if keyword == "产品定价":
        return f"{base}，价格背后真正决定了什么？"
    if keyword == "产品设计":
        return f"{base}，为什么用户真正买的不是产品？"
    if keyword == "利润模型":
        return f"{base}，为什么利润比规模更重要？"
    if keyword == "增长策略":
        return f"{base}，增长真正卡在哪里？"
    if keyword in {"自媒体起号", "涨粉方法"}:
        return f"{base}，新人最该先看懂什么？"
    if keyword == "内容定位":
        return f"{base}，定位错在哪里？"
    if keyword == "IP变现":
        return f"{base}，到底靠什么成交？"
    if keyword == "账号避坑":
        return f"{base}，最容易踩的坑是什么？"
    if keyword in {"AI应用", "AI提效", "AI团队", "AI就业", "模型工具", "自动化工作流", "AI认知", "AI付费", "豆包应用"}:
        return f"{base}，真正有效的用法是什么？"
    if keyword in {"副业赚钱", "第一桶金", "赚钱思维", "原始积累"}:
        return f"{base}，真正的关键是什么？"
    if keyword in {"房产选择", "资产配置", "金融认知", "财富认知"}:
        return f"{base}，怎么判断才不吃亏？"
    if keyword in {"职场成长", "工作选择", "求职选择", "执行力", "认知成长", "时间管理", "人生翻盘", "社交关系", "个人表达", "机会识别", "规则思维", "工作复利", "演讲表达", "阶层跃迁", "社群避坑", "资源交换"}:
        return f"{base}，真正的分水岭是什么？"
    if keyword in {"学习方法", "读书方法", "目标管理", "自我修炼"}:
        return f"{base}，用3步真正落到行动里？"
    if keyword in {"内容拍摄", "内容传播", "内容定位", "平台认知", "自媒趋势", "带货视频"}:
        return f"{base}，为什么多数人做反了？"
    if keyword in {"城市生活", "宠物旅行", "居住体验", "读书推荐", "线下活动", "人物访谈", "生活方式", "生活观察", "综合观察"}:
        return ensure_enhanced_when_no_element(base, base, keyword, elements)
    return ensure_enhanced_when_no_element(base, base, keyword, elements)


def parse_csv_table(path: Path, primary_bucket: str) -> Dict[Tuple[str, str], TopicRow]:
    """读取xlsx选题表，兼容旧列数格式。函数名保留兼容旧调用。"""
    rows: Dict[Tuple[str, str], TopicRow] = {}
    if not path.exists():
        return rows
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        header_raw = next(iterator, None)
        if not header_raw:
            return rows
        header = [normalize_cell(h) for h in header_raw]
    except Exception:
        return rows
    is_legacy_8col = header == ["核心关键词", "原选题", "原爆款元素", "博主名", "点赞数", "链接", "发布时间", "优化后爆款选题"]
    is_legacy_10col = header == ["核心关键词", "原选题", "原爆款元素", "博主名", "点赞数", "链接", "发布时间", "优化后爆款选题", "是否选中", "结构类型"]
    is_current = header == TABLE_HEADER
    is_v41 = header == LEGACY_V41_HEADER
    if not is_current and not is_v41 and not is_legacy_8col and not is_legacy_10col:
        return rows
    for cells_raw in iterator:
        cells = [normalize_cell(v) for v in cells_raw]
        if is_current or is_v41:
            if len(cells) < len(TABLE_HEADER):
                cells = cells + [""] * (len(TABLE_HEADER) - len(cells))
            row = TopicRow(
                keyword=cells[0],
                original_topic=cells[1],
                original_elements=cells[2],
                primary_bucket=primary_bucket,
                blogger=cells[3],
                likes=cells[4],
                link=cells[5],
                optimized_topic=cells[1],
                is_selected=cells[6] if len(cells) > 6 else "",
                benchmark_case_id=cells[7] if len(cells) > 7 else "",
                status=cells[8] if len(cells) > 8 else "",
            )
            rows[row.key()] = row
            continue
        if is_legacy_8col and len(cells) >= 8:
            row = TopicRow(
                keyword=cells[0],
                original_topic=cells[1],
                original_elements=cells[2],
                primary_bucket=primary_bucket,
                blogger=cells[3],
                likes=cells[4],
                link=cells[5],
                optimized_topic=cells[7],
                is_selected="",
                benchmark_case_id="",
                status="",
            )
            rows[row.key()] = row
            continue
        if is_legacy_10col and len(cells) >= 10:
            row = TopicRow(
                keyword=cells[0],
                original_topic=cells[1],
                original_elements=cells[2],
                primary_bucket=primary_bucket,
                blogger=cells[3],
                likes=cells[4],
                link=cells[5],
                optimized_topic=cells[7],
                is_selected=cells[8],
                benchmark_case_id="",
                status="",
            )
            rows[row.key()] = row
            continue
    return rows


# ── markdown 选题表解析（与仓库 .md 格式一致）──
def parse_markdown_table(path: Path, primary_bucket: str) -> Dict[Tuple[str, str], TopicRow]:
    """读取 markdown 选题表，解析为 TopicRow 字典。

    兼容三种表头：
    - 当前 V4.2：核心关键词 / 选题 / 原爆款元素 / 博主名 / 点赞数 / 链接 / 是否选中 / 对标复刻拆解编号 / 状态
    - 旧 10 列（含发布时间、优化后爆款选题）
    - 旧 8 列（含发布时间、优化后爆款选题，无是否选中/结构类型）
    按列名定位，不依赖固定列序。
    """
    rows: Dict[Tuple[str, str], TopicRow] = {}
    if not path.exists():
        return rows
    try:
        text = path.read_text(encoding="utf-8")
    except Exception:
        return rows

    lines = text.splitlines()
    header: List[str] = []
    data_lines: List[str] = []
    for line in lines:
        s = line.strip()
        if not s.startswith("|"):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if not header:
            # 跳过分隔行（全为 - 或 :--）
            if all(re.match(r":?-{3,}:?", c) for c in cells if c):
                continue
            header = cells
            continue
        if all(re.match(r":?-{3,}:?", c) for c in cells if c):
            continue
        data_lines.append(line)

    if not header:
        return rows

    def idx(name: str) -> int:
        return header.index(name) if name in header else -1

    i_kw = idx("核心关键词")
    i_topic = idx("选题") if idx("选题") >= 0 else idx("原选题")
    i_elem = idx("原爆款元素")
    i_blogger = idx("博主名")
    i_likes = idx("点赞数")
    i_link = idx("链接")
    i_opt = idx("优化后爆款选题")
    i_sel = idx("是否选中")
    i_struct = idx("对标复刻拆解编号")
    i_status = idx("状态")

    def get(cells: List[str], i: int) -> str:
        return cells[i] if 0 <= i < len(cells) else ""

    for line in data_lines:
        s = line.strip().strip("|")
        cells = [c.strip() for c in s.split("|")]
        if i_topic < 0:
            continue
        topic = get(cells, i_topic)
        if not topic:
            continue
        row = TopicRow(
            keyword=get(cells, i_kw),
            original_topic=topic,
            original_elements=get(cells, i_elem),
            primary_bucket=primary_bucket,
            blogger=get(cells, i_blogger),
            likes=get(cells, i_likes),
            link=get(cells, i_link),
            optimized_topic=get(cells, i_opt),
            is_selected=get(cells, i_sel),
            benchmark_case_id=get(cells, i_struct),
            status=get(cells, i_status),
        )
        rows[row.key()] = row
    return rows


def _parse_likes(value: str) -> int:
    """把点赞数字符串解析为整数；空或非法按 0 处理。"""
    s = (value or "").strip().replace(",", "").replace("，", "")
    if not s:
        return 0
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return 0


def _row_is_better(cand: TopicRow, cur: TopicRow) -> bool:
    """cand 是否应替代 cur 成为保留行（用于选题去重）。

    优先级：点赞更高 > 已填写对标复刻拆解编号 > 是否选中=是 > 先出现（稳定）。
    """
    c_likes, cur_likes = _parse_likes(cand.likes), _parse_likes(cur.likes)
    if c_likes != cur_likes:
        return c_likes > cur_likes
    c_case, cur_case = bool(cand.benchmark_case_id.strip()), bool(cur.benchmark_case_id.strip())
    if c_case != cur_case:
        return c_case
    c_sel, cur_sel = cand.is_selected.strip() == "是", cur.is_selected.strip() == "是"
    if c_sel != cur_sel:
        return c_sel
    return False


def dedupe_by_topic(rows: List[TopicRow]) -> List[TopicRow]:
    """选题严格相等去重：同一 选题 文本只保留一行，保留点赞数最高的一行。

    规则（与用户确认 Q1/Q2）：
    - 严格相等：按 original_topic.strip() 完全相同判重，不做模糊归一化。
    - 保留点赞最高行；点赞相同则优先保留已填对标复刻拆解编号的行（数据更完整），
      其次保留「是否选中=是」的行，再其次保留先出现行（稳定）。
    - 不动 TopicRow.key() 的链接 / 博主+选题 去重逻辑。
    """
    seen: Dict[str, TopicRow] = {}
    order: List[str] = []
    for r in rows:
        key = r.original_topic.strip()
        if key not in seen:
            seen[key] = r
            order.append(key)
        elif _row_is_better(r, seen[key]):
            seen[key] = r
    return [seen[k] for k in order]


def write_category_table(path: Path, category: str, rows: List[TopicRow]) -> None:
    """将分类结果写为 markdown 选题表（与仓库 .md 格式一致）。

    格式：标题 + 统计块 + 居中对齐的 markdown 表格。
    统计块：待选中选题（是否选中为空）+ 待生成结构（已选+已填对标复刻拆解编号+状态为空）。
    """
    path.parent.mkdir(parents=True, exist_ok=True)

    # 统计
    pending_select = 0
    pending_structure = 0
    for r in rows:
        is_sel = r.is_selected.strip()
        struct = r.benchmark_case_id.strip()
        status = r.status.strip()
        if not is_sel:
            pending_select += 1
        if is_sel and struct and not status:
            pending_structure += 1

    # 计算每列显示宽度（中文按 2 宽）
    def disp_width(s: str) -> int:
        return sum(2 if ord(c) > 127 else 1 for c in s)

    n_col = len(TABLE_HEADER)
    widths = [disp_width(h) for h in TABLE_HEADER]
    for r in rows:
        cells = r.to_cells()
        for i in range(min(n_col, len(cells))):
            widths[i] = max(widths[i], disp_width(cells[i]))

    def fmt_row(cells: List[str]) -> str:
        out = []
        for i in range(n_col):
            c = markdown_escape(cells[i] if i < len(cells) else "")
            out.append(" " + c.ljust(widths[i]) + " ")
        return "|" + "|".join(out) + "|"

    lines = []
    lines.append(f"# {path.stem}")
    lines.append("")
    lines.append("**统计：**")
    lines.append(f"*1、待选中选题，共 {pending_select} 条*")
    lines.append(f"*2、待生成结构，共 {pending_structure} 条*")
    lines.append("*3、已生成结构，共 0 条*")
    lines.append("*4、已手动优化，共 0 条*")
    lines.append("*5、已生成正文，共 0 条*")
    lines.append("")
    lines.append(fmt_row(TABLE_HEADER))
    lines.append("|" + "|".join(" " + "-" * widths[i] + " " for i in range(n_col)) + "|")
    for r in rows:
        lines.append(fmt_row(r.to_cells()))
    lines.append("")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    append_brand_footer(path)


def detect_excel_format(path: Path) -> str:
    with path.open("rb") as file_obj:
        head = file_obj.read(8)
    if head.startswith(XLSX_SIGNATURE):
        return "xlsx"
    if head.startswith(XLS_SIGNATURE):
        return "xls"
    return "unknown"


def read_xlsx(path: Path) -> Tuple[List[str], List[List[str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = next(iterator, None)
    if not headers:
        return [], []
    normalized_headers = [normalize_cell(h) for h in headers]
    rows: List[List[str]] = []
    for row in iterator:
        rows.append([normalize_cell(v) for v in row])
    return normalized_headers, rows


def read_xls(path: Path) -> Tuple[List[str], List[List[str]]]:
    if xlrd is None:
        raise RuntimeError("缺少依赖 xlrd，无法读取老式 Excel。")
    workbook = xlrd.open_workbook(path)
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows <= 0:
        return [], []
    headers = [normalize_cell(value) for value in sheet.row_values(0)]
    rows: List[List[str]] = []
    for row_index in range(1, sheet.nrows):
        rows.append([normalize_cell(value) for value in sheet.row_values(row_index)])
    return headers, rows


def read_excel_rows(path: Path) -> Tuple[str, List[str], List[List[str]]]:
    file_format = detect_excel_format(path)
    if file_format == "xlsx":
        headers, rows = read_xlsx(path)
        return file_format, headers, rows
    if file_format == "xls":
        headers, rows = read_xls(path)
        return file_format, headers, rows
    raise RuntimeError("无法识别的 Excel 文件格式")


def iter_input_files(input_dir: Path, blogger: Optional[str]) -> List[Path]:
    files = sorted({*input_dir.glob("*.xlsx"), *input_dir.glob("*.xls")})
    if blogger:
        files = [p for p in files if p.stem == blogger]
    return files


# 已知非标题列：# 密度兜底时排除这些列，避免链接列（URL 可能含 #）等被误判为标题列
NON_TITLE_COLUMNS = {
    "链接", "点赞数", "点赞", "发布时间", "分类", "博主名", "博主",
    "备注", "标签", "时长", "播放量", "播放", "评论数", "收藏数", "转发数",
}


def resolve_title_column(headers: List[str], data_rows: List[List[str]]) -> Optional[str]:
    """解析标题列（即原「视频信息」列）。

    识别优先级：
    1. 列名精确匹配「视频信息」→ 直接返回（保住现有老文件，零回归）。
    2. 否则按「# 密度」兜底：统计每列含 # 的单元格数，排除已知非标题列后，
       取含 # 最多且严格大于其他列的列（标题列通常每条都带 #话题）。
    找不到返回 None（调用方据此报 missing_required，不瞎猜）。
    """
    if "视频信息" in headers:
        return "视频信息"
    candidates = [h for h in headers if h and h not in NON_TITLE_COLUMNS]
    if not candidates:
        return None
    hash_count: Dict[str, int] = {}
    for h in candidates:
        idx = headers.index(h)
        cnt = 0
        for cells in data_rows:
            if idx < len(cells) and "#" in str(cells[idx]):
                cnt += 1
        if cnt > 0:
            hash_count[h] = cnt
    if not hash_count:
        return None
    best = max(hash_count, key=lambda h: hash_count[h])
    others = [v for k, v in hash_count.items() if k != best]
    if others and hash_count[best] <= max(others):
        return None  # 与其它列并列，不瞎猜
    return best


def classify_file(path: Path, audit: Dict[str, List[str]], processed_registry: Optional[set] = None,
                   output_dir: Optional[Path] = None) -> List[TopicRow]:
    blogger = path.stem
    try:
        file_format, headers, data_rows = read_excel_rows(path)
    except Exception as exc:
        audit["unreadable"].append(f"{path.name}：{type(exc).__name__} {exc}")
        return []

    title_col = resolve_title_column(headers, data_rows)
    if not title_col:
        audit["missing_required"].append(
            f"{path.name}：未找到标题列（尝试列名「视频信息」与 # 密度识别均失败），请检查表头"
        )
        return []
    if title_col != "视频信息":
        audit["readable"].append(f"{path.name}：标题列经 # 自动识别为「{title_col}」")

    header_index = {name: i for i, name in enumerate(headers)}
    result: List[TopicRow] = []
    for offset, cells in enumerate(data_rows, start=2):
        def get(name: str) -> str:
            i = header_index.get(name)
            if i is None or i >= len(cells):
                return ""
            return cells[i]

        raw_topic = get(title_col)
        topic, shortened, classify_basis, shorten_reason = clean_topic(raw_topic)
        if not topic:
            audit["missing_topic"].append(f"{path.name} 第{offset}行：{title_col}为空")
            continue

        # ── 注册表跳过：仅在普通增量模式使用。恢复/重建模式必须重新从源表构造行。──
        row_link = get("链接")
        if processed_registry is not None:
            reg_key = ("link", row_link) if row_link else ("topic", f"{blogger}::{topic}")
            if reg_key in processed_registry:
                audit.setdefault("registry_skipped", []).append(
                    f"{path.name} 第{offset}行：已在处理注册表中，跳过 -> {topic}"
                )
                continue

        if shortened:
            audit["long_topic"].append(
                f"{path.name} 第{offset}行：原始标题={classify_basis} -> 提炼后={topic}；原因={shorten_reason or '标题过长，保留完整主张句'}"
            )

        initial_bucket, _ = classify_topic(topic)
        keyword = infer_core_keyword(topic, initial_bucket)
        primary_bucket = keyword_to_primary_bucket(keyword, initial_bucket)
        if primary_bucket == "其他类型":
            display_category, display_reason = infer_other_display_category(topic, blogger)
            suspected_category, suspected_reason = infer_suspected_category(topic)
            suspected_label = "无" if suspected_category == "待细化主题" else suspected_category
            suspected_flag = "否" if suspected_label == "无" else "是"
            audit["other_reason"].append(
                f"{path.name} 第{offset}行：进入 99 缓冲区 -> 选题={topic}；当前具体命名={display_category}；是否疑似应归入五大类={suspected_flag}；疑似应归类={suspected_label}；原因：{display_reason if suspected_flag == '否' else suspected_reason}"
            )
        original_elements = detect_original_elements(topic)
        row = TopicRow(
            keyword=keyword,
            original_topic=topic,
            original_elements=original_elements,
            primary_bucket=primary_bucket,
            blogger=blogger,
            likes=get("点赞数"),
            link=get("链接"),
            optimized_topic=optimize_topic(topic, keyword, original_elements),
        )
        if has_problematic_topic_residue(row.original_topic):
            audit["topic_residue"].append(
                f"{path.name} 第{offset}行：提炼后仍疑似长文案/提示词残留 -> {row.original_topic}"
            )
        if row.keyword == "待细化主题":
            audit["keyword_missing"].append(
                f"{path.name} 第{offset}行：有效选题仍缺核心关键词 -> {row.original_topic}"
            )
        if row.original_elements == "无明显元素" and row.original_topic == row.optimized_topic:
            audit["unoptimized_no_element"].append(
                f"{path.name} 第{offset}行：无明显元素但优化标题原样复制 -> {row.original_topic}"
            )
        if not row.link:
            audit["missing_link"].append(f"{path.name} 第{offset}行：缺少链接，用 博主名+选题 去重")
        result.append(row)

    # ── 爆款过滤：点赞数 > 1000 或 在文件内排前 20% ──
    def _parse_likes(r: TopicRow) -> int:
        try:
            return int(float(r.likes))
        except (ValueError, TypeError):
            return 0

    all_likes = sorted([_parse_likes(r) for r in result], reverse=True)
    top20_count = max(1, -(-len(all_likes) // 5))  # ceil(n*0.2), 至少保留1条
    top20_threshold = all_likes[top20_count - 1] if all_likes else 0

    kept: List[TopicRow] = []
    for r in result:
        likes_val = _parse_likes(r)
        if likes_val > 1000 or likes_val >= top20_threshold:
            kept.append(r)
        else:
            audit.setdefault("filtered_out", []).append(
                f"{path.name}：选题={r.original_topic}，点赞={r.likes}，原因=不满1000且不在前20%（阈值={top20_threshold}）"
            )
            # ── 爆款过滤剔除的选题也写入注册表，避免下次重复过滤 ──
            reg_key = ("link", r.link) if r.link else ("topic", f"{r.blogger}::{r.original_topic}")
            if output_dir is not None:
                append_registry_record(
                    output_dir, reg_key, r.blogger, r.original_topic,
                    "__filtered__", r.keyword, r.link, path.name,
                    "爆款过滤剔除，已注册避免重复过滤",
                )

    audit["readable"].append(
        f"{path.name}：按 {file_format} 读取 {len(data_rows)} 行，生成候选 {len(result)} 条，爆款过滤后保留 {len(kept)} 条"
    )
    return kept


def count_topic_rows(existing: Dict[str, Dict[Tuple[str, str], TopicRow]]) -> int:
    return sum(len(rows) for rows in existing.values())


def total_candidate_rows(existing: Dict[str, Dict[Tuple[str, str], TopicRow]]) -> int:
    return sum(len(rows) for rows in existing.values())


def topic_table_paths(output_dir: Path) -> List[Path]:
    paths = []
    for path in sorted(output_dir.glob(TOPIC_TABLE_GLOB)):
        if path.name == HAND_INPUT_FILENAME:
            continue
        paths.append(path)
    return paths


def load_existing(output_dir: Path) -> Dict[str, Dict[Tuple[str, str], TopicRow]]:
    existing: Dict[str, Dict[Tuple[str, str], TopicRow]] = {}
    for category, filename in CATEGORIES.items():
        existing[category] = parse_markdown_table(output_dir / filename, category)
    return existing


def find_existing_category(existing: Dict[str, Dict[Tuple[str, str], TopicRow]], key: Tuple[str, str]) -> Optional[str]:
    for category, rows in existing.items():
        if key in rows:
            return category
    return None


# ── 选题处理注册表（与选题表解耦，追加式记录，避免删除选题后重复分类）──
REGISTRY_FILENAME = ".processed_registry.jsonl"


def load_processed_registry(output_dir: Path) -> set:
    """加载选题处理注册表，返回已处理选题的去重键集合。

    注册表是追加式 JSONL 文件，记录所有已处理过的选题（含已分类和已爆款过滤剔除），
    与 6 张选题表的增删完全解耦。即使某条选题已从选题表中被用户删除，注册表中仍有记录，
    下次运行不会被当作新选题重复分类。

    去重键格式与 TopicRow.key() 完全一致：("link", url) 或 ("topic", f"{blogger}::{topic}")。
    """
    registry_path = output_dir / REGISTRY_FILENAME
    processed: set = set()
    if not registry_path.exists():
        return processed
    try:
        for line in registry_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
            except json.JSONDecodeError:
                continue
            key = obj.get("key")
            if isinstance(key, list) and len(key) == 2:
                processed.add((key[0], key[1]))
    except Exception:
        pass
    return processed


def append_registry_record(output_dir: Path, key: Tuple[str, str], blogger: str, topic: str,
                           category: str, keyword: str, link: str, source_file: str, note: str) -> None:
    """向注册表追加一条处理记录（追加式写入，不修改已有记录）。"""
    import datetime as _dt
    registry_path = output_dir / REGISTRY_FILENAME
    record = {
        "key": list(key),
        "source_type": "topic_table" if category else "input_excel",
        "blogger": blogger,
        "topic": topic,
        "category": category,
        "keyword": keyword,
        "link": link,
        "processed_at": _dt.datetime.now().isoformat(timespec="seconds"),
        "source_file": source_file,
        "note": note,
    }
    with registry_path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(record, ensure_ascii=False) + "\n")


def write_audit(audit_dir: Path, audit: Dict[str, List[str]], totals: Dict[str, int], output_dir: Path) -> Path:
    audit_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    path = audit_dir / f"{stamp}_爆款选题分类审核.md"
    lines = [
        "# 爆款选题分类审核",
        "",
        f"- 输出目录：{output_dir}",
        f"- 写入新选题：{totals['inserted']}",
        f"- 更新旧选题：{totals.get('updated', 0)}",
        f"- 跳过重复：{totals['duplicates']}",
        f"- 移动分类：{totals.get('moved', 0)}",
        f"- 关键词修正：{totals.get('keyword_rewritten', 0)}",
        f"- 标题重写：{totals.get('title_rewritten', 0)}",
        f"- 长文案残留：{totals.get('topic_residue', 0)}",
        f"- 缺核心关键词：{totals.get('keyword_missing', 0)}",
        f"- 无元素原样复制：{totals.get('unoptimized_no_element', 0)}",
        f"- 本轮新进入其他类型：{totals['other']}",
        f"- 当前其他类型总量：{totals.get('other_total', totals['other'])}",
        f"- 爆款过滤剔除：{totals.get('filtered', 0)}",
        f"- 注册表跳过：{totals.get('registry_skipped', 0)}",
        f"- 选题去重合并：{totals.get('topic_deduped', 0)}",
        f"- 索引表字段：{' / '.join(TABLE_HEADER)}",
        "",
    ]
    sections = [
        ("可读取账号表", "readable"),
        ("不可读取账号表", "unreadable"),
        ("缺少必需字段", "missing_required"),
        ("缺少视频信息", "missing_topic"),
        ("缺少链接", "missing_link"),
        ("长选题提炼", "long_topic"),
        ("长文案残留复核", "topic_residue"),
        ("缺核心关键词复核", "keyword_missing"),
        ("无元素原样复制复核", "unoptimized_no_element"),
        ("迁移分类", "category_moved"),
        ("关键词修正", "keyword_rewritten"),
        ("标题重写", "title_rewritten"),
        ("财富/成长冲突样例", "wealth_growth_conflict"),
        ("其他类型原因", "other_reason"),
        ("爆款元素与关键词说明", "topic_index"),
        ("爆款过滤剔除", "filtered_out"),
        ("注册表跳过（已处理不再重复分类）", "registry_skipped"),
        ("恢复模式", "recovery"),
        ("选题去重合并", "topic_deduped"),
    ]
    for title, key in sections:
        lines.extend([f"## {title}", ""])
        items = audit.get(key, [])
        if not items:
            lines.append("- 无")
        else:
            lines.extend(f"- {item}" for item in items)
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    append_brand_footer(path)
    return path


def legacy_manual_table_state(output_dir: Path) -> Tuple[bool, Optional[float]]:
    manual = output_dir / HAND_INPUT_FILENAME
    if manual.exists():
        return True, manual.stat().st_mtime
    return False, None


def main() -> int:
    parser = argparse.ArgumentParser(description="从对标账号库分类生成爆款选题表")
    parser.add_argument("--root", default=None, help="工作区根目录，默认读取 AI_TRAFFIC_FACTORY_ROOT")
    parser.add_argument("--input-dir", default=None, help="账号表目录")
    parser.add_argument("--output-dir", default=None, help="爆款选题库输出目录")
    parser.add_argument("--audit-dir", default=None, help="审核报告目录")
    parser.add_argument("--blogger", default=None, help="只处理指定博主文件名，不含扩展名")
    parser.add_argument("--rebuild-from-source", action="store_true", help="忽略处理注册表跳过，从原始 Excel 重建六张选题表")
    args = parser.parse_args()

    script_path = Path(__file__).resolve()
    root_value = args.root or os.environ.get("AI_TRAFFIC_FACTORY_ROOT")
    root = Path(root_value).resolve() if root_value else script_path.parents[3]
    input_dir = Path(args.input_dir).resolve() if args.input_dir else root / "02_资产中心" / "04_选题库/01_对标账号"
    output_dir = Path(args.output_dir).resolve() if args.output_dir else root / "02_资产中心" / "04_选题库/02_选题分类"
    audit_dir = (
        Path(args.audit_dir).resolve()
        if args.audit_dir
        else root / ".workbuddy" / "local_only" / "reports" / "dry-goods-pipeline"
    )
    if not input_dir.exists():
        raise SystemExit(f"输入目录不存在：{input_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)

    legacy_manual_exists_before, legacy_manual_mtime_before = legacy_manual_table_state(output_dir)
    previous_existing = load_existing(output_dir)
    processed_registry = load_processed_registry(output_dir)
    previous_row_count = count_topic_rows(previous_existing)
    input_files = iter_input_files(input_dir, args.blogger)
    recovery_mode = args.rebuild_from_source or (previous_row_count == 0 and bool(processed_registry) and bool(input_files))
    registry_for_skip = None if recovery_mode else processed_registry
    existing: Dict[str, Dict[Tuple[str, str], TopicRow]] = {category: {} for category in CATEGORIES}
    audit: Dict[str, List[str]] = {
        "readable": [],
        "unreadable": [],
        "missing_required": [],
        "missing_topic": [],
        "missing_link": [],
        "long_topic": [],
        "topic_residue": [],
        "keyword_missing": [],
        "unoptimized_no_element": [],
        "other_reason": [],
        "category_moved": [],
        "keyword_rewritten": [],
        "title_rewritten": [],
        "wealth_growth_conflict": [],
        "topic_index": [],
        "filtered_out": [],
        "registry_skipped": [],
        "recovery": [],
        "topic_deduped": [],
    }
    if recovery_mode:
        reason = "--rebuild-from-source" if args.rebuild_from_source else "选题表为空且处理注册表非空，自动进入恢复模式"
        audit["recovery"].append(
            f"{reason}；注册表记录数={len(processed_registry)}；恢复前表内行数={previous_row_count}"
        )
    totals = {
        "inserted": 0,
        "updated": 0,
        "duplicates": 0,
        "other": 0,
        "moved": 0,
        "keyword_rewritten": 0,
        "title_rewritten": 0,
        "topic_residue": 0,
        "keyword_missing": 0,
        "unoptimized_no_element": 0,
        "filtered": 0,
        "registry_skipped": 0,
        "topic_deduped": 0,
    }

    for file_path in input_files:
        for row in classify_file(file_path, audit, registry_for_skip, output_dir):
            key = row.key()
            existing_category = find_existing_category(previous_existing, key)
            current = previous_existing.get(existing_category, {}).get(key) if existing_category else None
            if current:
                row.is_selected = current.is_selected
                row.benchmark_case_id = current.benchmark_case_id
                row.status = current.status
            bucket = existing.setdefault(row.primary_bucket, {})
            if key in bucket:
                totals["duplicates"] += 1
                continue
            if existing_category == row.primary_bucket:
                if current and current.to_cells() == row.to_cells():
                    totals["duplicates"] += 1
                else:
                    if current and current.keyword != row.keyword:
                        totals["keyword_rewritten"] += 1
                        audit["keyword_rewritten"].append(
                            f"{row.blogger}：{row.original_topic}；{current.keyword} -> {row.keyword}"
                        )
                    if current and current.optimized_topic != row.optimized_topic:
                        totals["title_rewritten"] += 1
                        audit["title_rewritten"].append(
                            f"{row.blogger}：{row.original_topic}；旧={current.optimized_topic}；新={row.optimized_topic}"
                        )
                    totals["updated"] += 1
                bucket[key] = row
                continue
            if existing_category and existing_category != row.primary_bucket:
                old_row = current
                audit["category_moved"].append(
                    f"{row.blogger}：{row.original_topic}；{existing_category} -> {row.primary_bucket}；关键词={row.keyword}"
                )
                if old_row and old_row.keyword != row.keyword:
                    totals["keyword_rewritten"] += 1
                    audit["keyword_rewritten"].append(
                        f"{row.blogger}：{row.original_topic}；{old_row.keyword} -> {row.keyword}"
                    )
                if old_row and old_row.optimized_topic != row.optimized_topic:
                    totals["title_rewritten"] += 1
                    audit["title_rewritten"].append(
                        f"{row.blogger}：{row.original_topic}；旧={old_row.optimized_topic}；新={row.optimized_topic}"
                    )
                if existing_category == "能力成长" and row.primary_bucket == "赚钱财富":
                    audit["wealth_growth_conflict"].append(
                        f"{row.blogger}：{row.original_topic}；原分类=能力成长；新分类=赚钱财富；原因=标题存在明确赚钱/收入/生意/资产承诺，财富结果优先于成长词"
                    )
                totals["moved"] += 1
            bucket[key] = row
            if not existing_category:
                totals["inserted"] += 1
                # ── 新选题插入时追加到处理注册表（内存+文件），永久记录避免重复分类 ──
                if key not in processed_registry:
                    processed_registry.add(key)
                    append_registry_record(
                        output_dir, key, row.blogger, row.original_topic,
                        row.primary_bucket, row.keyword, row.link, "",
                        "本轮新分类进入选题表",
                    )
            if row.primary_bucket == "其他类型":
                totals["other"] += 1

    candidate_row_count = total_candidate_rows(existing)
    if input_files and processed_registry and candidate_row_count == 0:
        raise SystemExit(
            "保护中止：输入库和处理注册表均存在，但本轮准备写入 0 条选题；"
            "为避免再次覆盖六张正式表，已停止写表。需要恢复时请使用 --rebuild-from-source。"
        )

    for category, filename in CATEGORIES.items():
        rows = sorted(existing.get(category, {}).values(), key=lambda item: item.keyword)
        # V4: 过滤「是否选中=否」的行
        rows = [r for r in rows if r.is_selected.strip() != "否"]
        # V4.2: 选题严格相等去重，保留点赞最高一行
        before = len(rows)
        rows = dedupe_by_topic(rows)
        merged = before - len(rows)
        if merged:
            audit["topic_deduped"].append(f"{category}：选题重复合并 {merged} 行")
            totals["topic_deduped"] += merged
        write_category_table(output_dir / filename, category, rows)

    totals["other_total"] = len(existing.get("其他类型", {}))
    totals["filtered"] = len(audit.get("filtered_out", []))
    totals["registry_skipped"] = len(audit.get("registry_skipped", []))
    for category_rows in existing.values():
        for row in category_rows.values():
            if has_problematic_topic_residue(row.original_topic):
                totals["topic_residue"] += 1
            if row.keyword == "待细化主题":
                totals["keyword_missing"] += 1
            if row.original_elements == "无明显元素" and row.original_topic == row.optimized_topic:
                totals["unoptimized_no_element"] += 1
    audit["topic_index"].append("已按 成本 / 人群 / 奇葩 / 难吃难用差 / 反向操作 / 怀旧 / 荷尔蒙 / 头牌 判断原爆款元素，并生成核心关键词（V4 已移除「优化后爆款选题」列，不再输出）。")

    legacy_manual_exists_after, legacy_manual_mtime_after = legacy_manual_table_state(output_dir)
    if legacy_manual_exists_before and legacy_manual_exists_after and legacy_manual_mtime_before != legacy_manual_mtime_after:
        raise SystemExit("保护失败：旧入口 00_手动输入选题表.xlsx 不应被直接改写")

    audit_path = write_audit(audit_dir, audit, totals, output_dir)
    print({
        "inserted": totals["inserted"],
        "duplicates": totals["duplicates"],
        "moved": totals["moved"],
        "keyword_rewritten": totals["keyword_rewritten"],
        "title_rewritten": totals["title_rewritten"],
        "topic_residue": totals["topic_residue"],
        "keyword_missing": totals["keyword_missing"],
        "unoptimized_no_element": totals["unoptimized_no_element"],
        "other": totals["other"],
        "filtered": totals["filtered"],
        "audit_path": str(audit_path),
    })
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
